import io
import unittest
from unittest.mock import patch, AsyncMock

from fastapi import HTTPException, UploadFile
from fastapi import BackgroundTasks
from PIL import Image

from classroom_app.routers.ai import _prepare_chat_uploads, _extract_exam_source_items, _build_exam_image_inputs, _EXAM_SOURCE_MAX_IMAGES
from classroom_app.services.ai_model_policy import resolve_execution_plan


def png_bytes():
    image = Image.new('RGB', (200, 100), 'white')
    target = io.BytesIO()
    image.save(target, format='PNG')
    return target.getvalue()


def upload(name, content):
    return UploadFile(file=io.BytesIO(content), filename=name)


class ChatUploadManifestTests(unittest.IsolatedAsyncioTestCase):
    async def test_oversized_exam_rejected_before_paper_or_job_creation(self):
        from classroom_app.routers.ai import ai_generate_exam
        data = {'title': 'Oversized exam', 'scope': 'Create questions about the supplied classroom topic', 'total_questions': 41}
        with patch('classroom_app.routers.ai._parse_exam_generation_request', new=AsyncMock(return_value=(data, []))), patch('classroom_app.routers.ai.get_db_connection') as database, patch('classroom_app.routers.ai.persist_ai_job_artifact') as artifact:
            with self.assertRaises(HTTPException) as raised:
                await ai_generate_exam(None, BackgroundTasks(), {'id': 1, 'role': 'teacher'})
            self.assertEqual(raised.exception.status_code, 400)
            database.assert_not_called()
            artifact.assert_not_called()

    async def test_exam_references_reject_empty_and_too_many_images(self):
        with self.assertRaises(HTTPException):
            await _extract_exam_source_items([upload('empty.txt', b'')], [], {'id':1, 'role':'teacher'})
        with self.assertRaises(HTTPException) as raised:
            _build_exam_image_inputs([{'type':'image', 'name':str(i), 'data_url':'data:image/png;base64,a'} for i in range(_EXAM_SOURCE_MAX_IMAGES + 1)])
        self.assertEqual(raised.exception.status_code, 413)
        manifest = await _extract_exam_source_items([upload('source.png', png_bytes())], [], {'id':1, 'role':'teacher'})
        self.assertEqual(len(_build_exam_image_inputs(manifest)), 1)

    def test_chat_context_keeps_visual_lite_and_plain_text_deepseek(self):
        context = {"operation": "chat", "source_feature": "classroom_chat", "class_offering_id": 1}
        for task in ('vision_interactive', 'deep_multimodal_reasoning'):
            plan = resolve_execution_plan(task, 'vision', context, environ={})
            self.assertEqual(plan.profile_id, 'vision_edge_low')
        plan = resolve_execution_plan('deep_text_reasoning', 'thinking', context, environ={})
        self.assertEqual(plan.provider, 'deepseek')

    async def test_text_before_image_keeps_correct_name_and_current_image(self):
        manifest = await _prepare_chat_uploads([upload('question.txt', b'question'), upload('answer.png', png_bytes())])
        self.assertEqual(manifest['file_texts'], [{'name': 'question.txt', 'content': 'question'}])
        self.assertEqual([item['name'] for item in manifest['image_inputs']], ['answer.png'])
        self.assertEqual(manifest['base64_urls'], [manifest['image_inputs'][0]['url']])
        self.assertNotIn('base64', str(manifest['attachments']))

    async def test_docx_text_and_embedded_image_both_reach_manifest(self):
        from docx import Document
        document = Document()
        document.add_paragraph('Question and screenshot below')
        document.add_picture(io.BytesIO(png_bytes()))
        stream = io.BytesIO()
        document.save(stream)
        manifest = await _prepare_chat_uploads([upload('answer.docx', stream.getvalue())])
        self.assertIn('Question and screenshot', manifest['file_texts'][0]['content'])
        self.assertEqual(len(manifest['image_inputs']), 1)
        self.assertTrue(manifest['image_inputs'][0]['name'].startswith('answer.docx / '))

    async def test_text_pdf_uses_text_but_scanned_and_vector_pdf_use_page_image(self):
        import fitz
        for kind in ('text', 'scan', 'vector'):
            with self.subTest(kind=kind), fitz.open() as doc:
                page = doc.new_page()
                if kind == 'scan':
                    page.insert_image(fitz.Rect(40, 40, 400, 220), stream=png_bytes())
                else:
                    page.insert_text((40, 40), 'Visible problem and instructions')
                    if kind == 'vector':
                        page.draw_rect(fitz.Rect(40, 90, 200, 250), color=(1, 0, 0))
                manifest = await _prepare_chat_uploads([upload('source.pdf', doc.tobytes())])
                self.assertEqual(len(manifest['image_inputs']), 0 if kind == 'text' else 1)

    async def test_broken_image_aborts_instead_of_silently_sending_other_files(self):
        with self.assertRaises(HTTPException) as raised:
            await _prepare_chat_uploads([upload('notes.txt', b'valid'), upload('broken.png', b'broken data')])
        self.assertEqual(raised.exception.status_code, 400)

    async def test_pixel_limit_checked_before_decode(self):
        with patch('PIL.Image.open') as opener:
            source = opener.return_value.__enter__.return_value
            source.width, source.height = 10_000, 10_000
            with self.assertRaises(HTTPException) as raised:
                await _prepare_chat_uploads([upload('huge.png', b'header')])
            self.assertEqual(raised.exception.status_code, 413)
            source.load.assert_not_called()

    async def test_scan_limit_and_empty_attachment_block(self):
        import fitz
        with fitz.open() as doc:
            for _ in range(9):
                doc.new_page()
            with self.assertRaises(HTTPException) as raised:
                await _prepare_chat_uploads([upload('too-many-pages.pdf', doc.tobytes())])
            self.assertEqual(raised.exception.status_code, 413)
        with self.assertRaises(HTTPException):
            await _prepare_chat_uploads([upload('empty.txt', b'')])

    async def test_document_image_overflow_reports_loss(self):
        import zipfile
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, 'w') as archive:
            archive.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:p><w:t>text</w:t></w:p></w:document>')
            for index in range(11):
                archive.writestr(f'word/media/image{index}.png', png_bytes())
        with self.assertRaises(HTTPException) as raised:
            await _prepare_chat_uploads([upload('many.docx', stream.getvalue())])
        self.assertIn('图片不完整', str(raised.exception.detail))


if __name__ == '__main__':
    unittest.main()
