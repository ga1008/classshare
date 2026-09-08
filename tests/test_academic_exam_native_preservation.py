from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import hashlib
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from lxml import etree

from classroom_app.services import academic_exam_analysis_native_service as native
from classroom_app.services.academic_final_material_document_service import build_exam_analysis_docx
from classroom_app.services.academic_final_material_source_service import NativeAcademicSourceError


def canonical(element):
    return etree.tostring(element, method="c14n")


class NativeExamPreservationTests(unittest.TestCase):
    def test_source_geometry_and_every_fixed_cell_are_preserved(self):
        before = Document(native._TEMPLATE_PATH)
        after = Document(BytesIO(build_exam_analysis_docx({
            "fields": {"course_name": "测试课程", "department_review_opinion": "已核"},
            "structured": {"analysis_text": "一、成绩分析\n教学内容。"},
        })))
        old, new = before.tables[0], after.tables[0]
        self.assertEqual(canonical(before.sections[0]._sectPr), canonical(after.sections[0]._sectPr))
        self.assertEqual(canonical(old._tbl.tblPr), canonical(new._tbl.tblPr))
        self.assertEqual(canonical(old._tbl.tblGrid), canonical(new._tbl.tblGrid))
        self.assertEqual(len(old.rows), len(new.rows))
        for source_row, result_row in zip(old.rows, new.rows):
            self.assertEqual(canonical(source_row._tr.trPr), canonical(result_row._tr.trPr))
            self.assertEqual(len(source_row._tr.tc_lst), len(result_row._tr.tc_lst))
            for source_cell, result_cell in zip(source_row._tr.tc_lst, result_row._tr.tc_lst):
                self.assertEqual(canonical(source_cell.tcPr), canonical(result_cell.tcPr))
        for row, column in ((0, 0), (3, 0), (7, 0), (9, 1), (13, 3), (17, 1),
                            (19, 0), (19, 1), (21, 0), (21, 1), (22, 0)):
            self.assertEqual(canonical(old.rows[row]._tr.tc_lst[column]),
                             canonical(new.rows[row]._tr.tc_lst[column]))

    def test_styles_settings_labels_and_other_untouched_parts_are_byte_identical(self):
        result = build_exam_analysis_docx({"fields": {"course_name": "测试课程"}, "structured": {}})
        with ZipFile(native._TEMPLATE_PATH) as source, ZipFile(BytesIO(result)) as output:
            for part in source.namelist():
                if part not in {"word/document.xml", "word/_rels/document.xml.rels",
                                "[Content_Types].xml", "word/media/image2.png"}:
                    self.assertEqual(source.read(part), output.read(part), part)
        before = Document(native._TEMPLATE_PATH)
        after = Document(BytesIO(result))
        self.assertEqual([canonical(node) for node in before._element.xpath(".//w:pict")],
                         [canonical(node) for node in after._element.xpath(".//w:pict")])

    def test_corrupted_template_never_rebuilds_an_approximation(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "template.docx"
            path.write_bytes(b"invalid")
            with patch.object(native, "_TEMPLATE_PATH", path), self.assertRaisesRegex(NativeAcademicSourceError, "模板校验失败"):
                build_exam_analysis_docx({})

    def test_custom_four_line_opinion_uses_top_of_slot_above_signature(self):
        document = Document(BytesIO(build_exam_analysis_docx({"fields": {
            "department_review_opinion": "审核意见" * 20,
        }})))
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
        cell = native._cell(document.tables[0], 20, 0)
        self.assertEqual(WD_CELL_VERTICAL_ALIGNMENT.TOP, cell.vertical_alignment)
        self.assertEqual(93, document.tables[0].rows[20].height.pt)
        self.assertEqual("审核意见" * 20, cell.text.replace("\n", ""))

    def test_original_source_survives_unwrapped_export_payload_and_reuses_its_chart(self):
        from classroom_app.services.academic_final_material_source_service import NativeAcademicSource, ACADEMIC_NATIVE_SOURCE_KEY
        from classroom_app.services.material_export_template_service import build_material_export_artifact

        content = b"verified-original"
        parsed = {"fields": {"course_name": "原件课程", "class_name": "原件班级"},
                  "distribution": [{"count": count, "ratio": 0} for count in (0, 1, 1, 1, 32)],
                  "statistics": {"average": 94.54}, "analysis_text": ""}
        with ZipFile(native._TEMPLATE_PATH) as package:
            images = {name: package.read(f"word/media/image{index}.png") for index, name in
                      ((1, "vertical_label_png"), (2, "chart_png"), (3, "analysis_label_png"))}
        images["hidden_text"] = " 32 "
        payload = {"template_key": "academic_exam_analysis", "fields": {"course_name": "旧解析课程"},
                   "structured": {"analysis_text": "教师填写的正文。"},
                   ACADEMIC_NATIVE_SOURCE_KEY: NativeAcademicSource(content, hashlib.sha256(content).hexdigest())}
        with patch("classroom_app.services.academic_exam_analysis_rtf_service.inspect_native_exam_analysis_source", return_value=images) as inspect_source, \
             patch("classroom_app.services.academic_final_material_service.parse_exam_analysis_rtf", return_value=parsed):
            artifact = build_material_export_artifact(payload, fallback_filename="analysis", requested_format="docx")
            inspect_source.assert_called_once_with(content)
            payload["structured"]["statistics"] = {"average": 10}
            with self.assertRaisesRegex(NativeAcademicSourceError, "成绩统计与教务原件不一致"):
                build_material_export_artifact(payload, fallback_filename="analysis", requested_format="docx")
        with ZipFile(BytesIO(artifact.content)) as output:
            self.assertEqual(images["chart_png"], output.read("word/media/image2.png"))
        document = Document(BytesIO(artifact.content))
        self.assertEqual("原件课程", native._cell(document.tables[0], 3, 1).text)
        self.assertEqual(" 32 ", native._cell(document.tables[0], 2, 0).text)
        self.assertEqual("教师填写的正文。", native._cell(document.tables[0], 18, 1).text)


if __name__ == "__main__":
    unittest.main()
