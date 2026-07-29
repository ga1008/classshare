"""过程材料命名（学年学期 + 课程 + 班级）的回归测试。

核心诉求：教师在材料库里必须能一眼分清"哪个班、哪门课、哪个学期"的材料——
同一门课有多个平行教学班，跨学期还会反复生成同类型材料。
"""

import unittest

from classroom_app.services.material_identity_service import (
    build_final_material_export_filename,
    build_final_material_package_name,
    context_summary,
    period_label,
)


FULL = {
    "academic_year": "2025-2026",
    "semester": "第二学期",
    "course_name": "动态web程序设计",
    "class_name": "软工2401班",
}
LABEL = "机试（作品设计）考核登分表"


class PeriodLabelTests(unittest.TestCase):
    def test_full_style_reads_year_and_term(self):
        self.assertEqual("2025-2026学年第二学期", period_label(FULL))

    def test_compact_style_matches_official_file_prefix(self):
        self.assertEqual("2025-2026-2", period_label(FULL, style="compact"))

    def test_first_term_is_distinguished(self):
        fields = {**FULL, "semester": "第一学期"}
        self.assertEqual("2025-2026学年第一学期", period_label(fields))
        self.assertEqual("2025-2026-1", period_label(fields, style="compact"))

    def test_canonical_single_string_is_accepted(self):
        # semester_identity_service 的 canonical 形式，学年学期写在一个字段里
        self.assertEqual("2025-2026学年第二学期", period_label({"semester": "2025-2026第二学期"}))

    def test_year_only_degrades_instead_of_guessing_term(self):
        self.assertEqual("2025-2026学年", period_label({"academic_year": "2025-2026"}))
        self.assertEqual("2025-2026", period_label({"academic_year": "2025-2026"}, style="compact"))

    def test_missing_period_returns_empty(self):
        self.assertEqual("", period_label({}))
        self.assertEqual("", period_label(None))
        self.assertEqual("", period_label({"academic_year": "未设置", "semester": "未设置"}))


class ContextSummaryTests(unittest.TestCase):
    def test_summary_joins_period_course_and_class(self):
        self.assertEqual(
            "2025-2026学年第二学期 · 动态web程序设计 · 软工2401班",
            context_summary(FULL),
        )

    def test_summary_skips_missing_parts_without_placeholders(self):
        self.assertEqual("动态web程序设计", context_summary({"course_name": "动态web程序设计"}))
        self.assertEqual("", context_summary({}))


class PackageNameTests(unittest.TestCase):
    def test_folder_name_carries_course_class_and_period(self):
        name = build_final_material_package_name(document_type_label=LABEL, fields=FULL)
        self.assertEqual(
            "AI生成-机试（作品设计）考核登分表-动态web程序设计-软工2401班-2025-2026学年第二学期",
            name,
        )

    def test_folder_name_omits_missing_segments_rather_than_padding(self):
        name = build_final_material_package_name(
            document_type_label=LABEL,
            fields={"course_name": "动态web程序设计"},
        )
        self.assertEqual("AI生成-机试（作品设计）考核登分表-动态web程序设计", name)
        self.assertNotIn("未命名", name)

    def test_folder_name_survives_empty_fields(self):
        name = build_final_material_package_name(document_type_label=LABEL, fields={})
        self.assertEqual("AI生成-机试（作品设计）考核登分表", name)

    def test_folder_name_strips_path_hostile_characters(self):
        name = build_final_material_package_name(
            document_type_label=LABEL,
            fields={**FULL, "class_name": "软工/2401:班"},
        )
        for char in '\\/:*?"<>|':
            self.assertNotIn(char, name)

    def test_parallel_classes_get_distinct_folder_names(self):
        first = build_final_material_package_name(document_type_label=LABEL, fields=FULL)
        second = build_final_material_package_name(
            document_type_label=LABEL, fields={**FULL, "class_name": "软工2402班"}
        )
        self.assertNotEqual(first, second)


class ExportFilenameTests(unittest.TestCase):
    def test_filename_matches_official_period_course_class_shape(self):
        self.assertEqual(
            "2025-2026-2《动态web程序设计》机试（作品设计）考核登分表-软工2401班.xlsx",
            build_final_material_export_filename(document_type_label=LABEL, fields=FULL),
        )

    def test_filename_fills_placeholders_so_gaps_are_visible(self):
        name = build_final_material_export_filename(document_type_label=LABEL, fields={})
        self.assertIn("未设学年学期", name)
        self.assertIn("未命名课程", name)
        self.assertIn("未命名班级", name)

    def test_filename_supports_official_sequence_prefix(self):
        name = build_final_material_export_filename(
            document_type_label=LABEL, fields=FULL, sequence=7
        )
        self.assertTrue(name.startswith("7. 2025-2026-2《"))

    def test_filename_normalizes_suffix(self):
        self.assertTrue(
            build_final_material_export_filename(
                document_type_label=LABEL, fields=FULL, suffix="docx"
            ).endswith(".docx")
        )

    def test_filename_has_no_illegal_characters(self):
        name = build_final_material_export_filename(
            document_type_label=LABEL,
            fields={**FULL, "course_name": "A/B:C*课"},
        )
        stem = name[:-5]
        for char in '\\/:*?"<>|':
            self.assertNotIn(char, stem)


class ExamGradeRecordIntegrationTests(unittest.TestCase):
    def test_normalize_sets_complete_export_filename(self):
        from classroom_app.services.exam_grade_record_service import (
            normalize_exam_grade_record_payload,
        )

        payload = normalize_exam_grade_record_payload(
            metadata={**FULL, "teacher_name": "张老师"}, tables=[], export_payload={}
        )
        filename = payload["fields"]["export_filename"]
        self.assertIn("2025-2026-2", filename)
        self.assertIn("动态web程序设计", filename)
        self.assertIn("软工2401班", filename)
        self.assertTrue(filename.endswith(".xlsx"))

    def test_exported_sheet_header_states_the_period(self):
        import io

        from openpyxl import load_workbook

        from classroom_app.services.exam_grade_record_service import (
            build_exam_grade_record_xlsx,
            normalize_exam_grade_record_payload,
        )

        payload = normalize_exam_grade_record_payload(
            metadata={**FULL, "teacher_name": "张老师"}, tables=[], export_payload={}
        )
        worksheet = load_workbook(io.BytesIO(build_exam_grade_record_xlsx(payload))).active
        header = str(worksheet["A2"].value or "")
        self.assertIn("课程：动态web程序设计", header)
        self.assertIn("专业年级班级：软工2401班", header)
        self.assertIn("学年学期：2025-2026学年第二学期", header)


if __name__ == "__main__":
    unittest.main()
