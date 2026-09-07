from __future__ import annotations

import asyncio
import json
import sqlite3
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from pydantic import ValidationError
from fastapi import HTTPException
from docx import Document
from starlette.requests import Request

from classroom_app.services import academic_final_material_service as service


class AcademicReviewOpinionTests(unittest.TestCase):
    def test_explicit_clear_is_not_refilled_from_older_metadata(self) -> None:
        payload = service.normalize_academic_final_material_payload(
            document_type="academic_exam_analysis",
            metadata={"department_review_opinion": "已核"},
            export_payload={"fields": {"department_review_opinion": ""}},
        )
        self.assertEqual("", payload["fields"]["department_review_opinion"])

    def setUp(self) -> None:
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.executescript("""
            CREATE TABLE signature_point_bindings (
                id INTEGER PRIMARY KEY, function_point_key TEXT, material_type TEXT,
                material_id TEXT, material_revision TEXT, signature_id INTEGER, display_order INTEGER
            );
            CREATE TABLE electronic_signatures (
                id INTEGER PRIMARY KEY, signature_kind TEXT, owner_role TEXT,
                subject_role TEXT, status TEXT, deleted_at TEXT
            );
            INSERT INTO electronic_signatures VALUES
                (1, 'personal', 'teacher', 'teacher', 'active', NULL),
                (2, 'personal', 'teacher', 'teacher', 'active', NULL),
                (3, 'stamp', 'system', 'other', 'active', NULL),
                (4, 'personal', 'teacher', 'teacher', 'active', NULL),
                (5, 'personal', 'teacher', 'teacher', 'inactive', NULL);
        """)
        self.record = {"id": 88, "teacher_id": 1, "document_type": service.ACADEMIC_EXAM_ANALYSIS_TYPE, "signature_revision": "current"}
        self.actor = patch.object(service.signature_service, "build_signature_actor", return_value={"id": 1, "role": "teacher"})
        self.actor_mock = self.actor.start()
        self.access = patch.object(service.signature_workflow_service, "signature_use_access_state", return_value={"can_use": True})
        self.access_mock = self.access.start()
        self.addCleanup(self.actor.stop)
        self.addCleanup(self.access.stop)
        self.payload = {
            "document_type": service.ACADEMIC_EXAM_ANALYSIS_TYPE,
            "fields": {key: "selected" for key in service.ACADEMIC_EXAM_ANALYSIS_EDIT_FIELDS},
            "structured": {"analysis_text": "教学分析"},
        }
        self.paths = patch.object(service, "resolve_signature_file_path", side_effect=lambda row: None if row["id"] == 4 else Path(f"/{row['id']}.png"))
        self.paths.start()
        self.compose = patch.object(service, "compose_signature_strip", side_effect=lambda paths, **kwargs: "|".join(paths))
        self.compose_mock = self.compose.start()
        self.addCleanup(self.paths.stop)
        self.addCleanup(self.compose.stop)
        self.addCleanup(self.conn.close)

    def bind(self, role: str, ids: list[int], *, record_id: int = 88, revision: str = "current") -> None:
        for order, signature_id in enumerate(ids):
            self.conn.execute(
                "INSERT INTO signature_point_bindings (function_point_key, material_type, material_id, material_revision, signature_id, display_order) "
                "VALUES (?, 'academic_final_material', ?, ?, ?, ?)",
                (f"academic_final_material.exam_analysis.{role}_review_signature", str(record_id), revision, signature_id, order),
            )

    def hydrate(self, *, include_images: bool = True) -> dict:
        return service.hydrate_academic_final_material_signature_paths(
            self.conn, self.payload, record=self.record, include_images=include_images,
        )["fields"]

    def test_signed_legacy_material_gets_role_specific_defaults_without_mutation(self) -> None:
        self.bind("department", [1])
        self.bind("dean", [2])
        fields = self.hydrate()
        self.assertEqual("已核", fields["department_review_opinion"])
        self.assertEqual("同意", fields["dean_review_opinion"])
        self.assertEqual("legacy_default", fields["department_review_opinion_source"])
        self.assertNotIn("department_review_opinion", self.payload["fields"])
        self.assertTrue(service.academic_exam_analysis_is_complete(fields, self.payload["structured"]))

    def test_unsigned_rebuilt_and_missing_files_do_not_gain_default_approval(self) -> None:
        self.payload["fields"].update({
            "department_signature_ids": [1], "department_signature_image_path": "/forged.png",
            "department_review_opinion_image_path": "/forged-stamp.png",
            "department_personal_signature_ids": [1], "department_review_stamp_ids": [3],
        })
        self.bind("department", [1], revision="old")
        self.bind("dean", [4, 5])
        fields = self.hydrate()
        self.assertEqual([], fields["department_personal_signature_ids"])
        self.assertEqual([], fields["dean_personal_signature_ids"])
        self.assertNotIn("department_review_opinion", fields)
        self.assertNotIn("dean_review_opinion", fields)
        self.assertNotIn("department_review_opinion_image_path", fields)
        self.assertNotIn("department_signature_image_path", fields)
        self.assertEqual("absent", fields["department_review_opinion_source"])
        self.assertFalse(service.academic_exam_analysis_is_complete(fields, self.payload["structured"]))

    def test_mixed_stamp_is_above_signature_and_does_not_inject_duplicate_text(self) -> None:
        self.bind("department", [3, 1])
        self.bind("dean", [2])
        fields = self.hydrate()
        self.assertEqual([3, 1], fields["department_signature_ids"])
        self.assertEqual([1], fields["department_personal_signature_ids"])
        self.assertEqual([3], fields["department_review_stamp_ids"])
        self.assertEqual(str(Path("/3.png")), fields["department_review_opinion_image_path"])
        self.assertEqual(str(Path("/1.png")), fields["department_signature_image_path"])
        self.assertNotIn("department_review_opinion", fields)
        self.assertEqual("bound_stamp", fields["department_review_opinion_source"])
        self.assertTrue(service.academic_exam_analysis_is_complete(fields, self.payload["structured"]))

    def test_explicit_custom_or_empty_remark_overrides_selected_stamp(self) -> None:
        self.bind("department", [3, 1])
        self.bind("dean", [2])
        self.payload["fields"]["department_review_opinion"] = "  已核\n 请改进\t分析  "
        fields = self.hydrate()
        self.assertEqual("已核 请改进 分析", fields["department_review_opinion"])
        self.assertNotIn("department_review_opinion_image_path", fields)
        self.payload["fields"]["department_review_opinion"] = ""
        fields = self.hydrate()
        self.assertEqual("", fields["department_review_opinion"])
        self.assertEqual("explicit", fields["department_review_opinion_source"])
        self.assertNotIn("department_review_opinion_image_path", fields)
        self.assertFalse(service.academic_exam_analysis_is_complete(fields, self.payload["structured"]))

    def test_stamp_alone_never_completes_review(self) -> None:
        self.bind("department", [3])
        self.bind("dean", [2])
        fields = self.hydrate()
        self.assertNotIn("department_signature_image_path", fields)
        self.assertFalse(service.academic_exam_analysis_is_complete(fields, self.payload["structured"]))

    def test_visibility_or_authorization_revoked_after_binding_removes_rendered_signature(self) -> None:
        self.bind("department", [1])
        self.payload["fields"]["department_signature_image_path"] = "/old-authorized.png"
        self.access_mock.return_value = {"can_use": False}
        fields = self.hydrate()
        self.assertEqual([], fields["department_signature_ids"])
        self.assertNotIn("department_signature_image_path", fields)
        self.assertNotIn("department_review_opinion", fields)
        self.actor_mock.assert_called_with(self.conn, {"role": "teacher", "id": 1})
        self.assertEqual("current", self.access_mock.call_args.kwargs["material_revision"])
        self.assertEqual("88", self.access_mock.call_args.kwargs["material_id"])

    def test_legacy_scalar_ids_and_image_free_editor_match_rendered_opinions(self) -> None:
        self.record["signature_revision"] = ""
        self.payload["fields"].update({"department_signature_id": 1, "dean_signature_id": 2})
        fields = self.hydrate(include_images=False)
        self.assertEqual("已核", fields["department_review_opinion"])
        self.assertEqual("同意", fields["dean_review_opinion"])
        self.assertNotIn("department_signature_image_path", fields)
        self.compose_mock.assert_not_called()

    def test_card_completion_is_recomputed_in_bulk_from_current_reviews(self) -> None:
        self.conn.executescript("""
            CREATE TABLE academic_final_material_batches (
                id TEXT, teacher_id INTEGER, analysis_record_id INTEGER, updated_at TEXT, edit_state_json TEXT
            );
            CREATE TABLE material_ai_import_records (
                id INTEGER, document_type_label TEXT, updated_at TEXT, signature_revision TEXT,
                content_quality_status TEXT, export_payload_json TEXT
            );
        """)
        for index in range(12):
            record_id = 88 + index
            payload = json.loads(json.dumps(self.payload))
            if index % 2:
                payload["fields"]["department_review_opinion"] = ""
            self.conn.execute("INSERT INTO academic_final_material_batches VALUES (?, 1, ?, '2026', ?)", (str(index), record_id, '{"analysis_complete": true}'))
            self.conn.execute("INSERT INTO material_ai_import_records VALUES (?, '试卷分析表', '2026', 'current', 'ok', ?)", (record_id, json.dumps(payload)))
            self.bind("department", [1], record_id=record_id)
            self.bind("dean", [2], record_id=record_id)
        statements = []
        self.conn.set_trace_callback(statements.append)
        with patch.object(service, "ensure_academic_final_material_schema"):
            items = service.list_teacher_final_material_batches(self.conn, 1, document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE)
        self.assertEqual(12, len(items))
        for item in items:
            self.assertEqual(int(item["id"]) % 2 == 0, item["edit_state"]["analysis_complete"])
        self.assertEqual(3, sum(sql.lstrip().upper().startswith("SELECT") for sql in statements))
        self.compose_mock.assert_not_called()

    def test_layout_metadata_matches_rendered_table_and_external_note(self) -> None:
        from classroom_app.services.academic_final_material_document_service import build_exam_analysis_docx

        payload = service.build_exam_analysis_export_payload({}, {})
        document = Document(BytesIO(build_exam_analysis_docx(payload)))
        self.assertEqual("gxufl-academic-exam-analysis-v3", payload["schema_version"])
        self.assertEqual(19, payload["layout_profile"]["table_rows"])
        self.assertEqual(len(document.tables[0].rows), payload["layout_profile"]["table_rows"])
        self.assertEqual("paragraph_after_table", payload["layout_profile"]["note_placement"])
        self.assertTrue(any(paragraph.text.startswith("注：1、") for paragraph in document.paragraphs))
        self.assertEqual(["heading", "opinion", "personal_signature_and_label"], payload["layout_profile"]["review_layout"]["content_order"])
        legacy = {"schema_version": "gxufl-academic-exam-analysis-v2", "fields": {"department_review_opinion": ""}, "layout_profile": {"table_rows": 22}}
        normalized = service.normalize_academic_final_material_payload(
            document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE, metadata={}, export_payload=legacy,
        )
        self.assertEqual(payload["layout_profile"], normalized["layout_profile"])
        self.assertEqual("", normalized["fields"]["department_review_opinion"])
        self.assertEqual(22, legacy["layout_profile"]["table_rows"])


class AcademicReviewOpinionApiTests(unittest.TestCase):
    def test_request_accepts_clear_and_rejects_oversized_remark(self) -> None:
        from classroom_app.routers.materials_parts import academic_final_materials as router

        body = router.AcademicFinalMaterialUpdateRequest(document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE, department_review_opinion="")
        self.assertEqual("", body.model_dump(exclude_unset=True)["department_review_opinion"])
        self.assertNotIn("dean_review_opinion", body.model_dump(exclude_unset=True))
        with self.assertRaises(ValidationError):
            router.AcademicFinalMaterialUpdateRequest(document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE, department_review_opinion="核" * 81)

    def test_first_personal_signature_defaults_but_explicit_clear_and_stamp_are_preserved(self) -> None:
        from classroom_app.routers.materials_parts import academic_final_materials as router

        for explicit, stamp, expected in ((False, False, "已核"), (True, False, ""), (False, True, None)):
            with self.subTest(explicit=explicit, stamp=stamp):
                fields = {"department_review_opinion": ""} if explicit else {}
                with (
                    patch.object(router.signature_service, "get_signature_row_for_actor", return_value=({"id": 1}, {})),
                    patch.object(router.signature_service, "resolve_signature_file_path", return_value=Path("/1.png")),
                    patch.object(router.signature_service, "is_stamp_signature", return_value=stamp),
                ):
                    intent = router._apply_signatures(
                        MagicMock(), {"id": 1}, fields, id_key="department_signature_id", ids_key="department_signature_ids",
                        path_key="department_signature_image_path", signature_ids=[1],
                        function_point_key="academic_final_material.exam_analysis.department_review_signature",
                        context_type="academic_final_material", context_id="88", context_label="试卷分析表",
                    )
                self.assertEqual(expected, fields.get("department_review_opinion"))
                self.assertEqual([1], intent["signature_ids"])

    def test_patch_persists_normalized_opinion_and_returns_latest_record(self) -> None:
        from classroom_app.routers.materials_parts import academic_final_materials as router

        record = {"id": 88, "parse_mode": "local_fallback", "export_payload_json": json.dumps({"document_type": service.ACADEMIC_EXAM_ANALYSIS_TYPE, "fields": {}, "structured": {}})}
        batch = {"edit_state_json": "{}", "class_offering_id": 3, "analysis_record_id": 88}
        latest = {"id": 88, "fields": {"department_review_opinion": "已核 请完善"}, "structured": {}, "preview_url": "/preview?v=current"}
        conn = MagicMock()
        connection_context = MagicMock()
        connection_context.__enter__.return_value = conn
        body = router.AcademicFinalMaterialUpdateRequest(document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE, department_review_opinion="  已核\n请完善  ")
        with (
            patch.object(router, "get_db_connection", return_value=connection_context),
            patch.object(router, "_batch_for_teacher", return_value=batch),
            patch.object(router, "_record_for_teacher", return_value=record),
            patch.object(router, "_persist_final_material_record_update", new_callable=AsyncMock, return_value={"id": 88}) as persist,
            patch.object(router, "_serialize_record_payload", return_value=latest),
            patch.object(router, "upsert_batch_state", return_value={"id": "batch"}),
        ):
            result = asyncio.run(router.api_update_academic_final_material(
                "batch", body, Request({"type": "http", "headers": []}), {"id": 1},
            ))
        saved = persist.call_args.args[2].export_payload
        self.assertEqual("已核 请完善", saved["fields"]["department_review_opinion"])
        self.assertEqual(latest, result["record"])
        self.assertTrue(persist.call_args.kwargs["require_unchanged_record"])

    def test_stale_editor_version_is_rejected_before_signatures_or_opinions_change(self) -> None:
        from classroom_app.routers.materials_parts import academic_final_materials as router

        connection_context = MagicMock()
        body = router.AcademicFinalMaterialUpdateRequest(
            document_type=service.ACADEMIC_EXAM_ANALYSIS_TYPE, department_review_opinion="已核",
            expected_updated_at="older-version",
        )
        with (
            patch.object(router, "get_db_connection", return_value=connection_context),
            patch.object(router, "_batch_for_teacher", return_value={"analysis_record_id": 88}),
            patch.object(router, "_record_for_teacher", return_value={"id": 88, "updated_at": "newer-version"}),
            patch.object(router, "_persist_final_material_record_update", new_callable=AsyncMock) as persist,
            self.assertRaises(HTTPException) as failure,
        ):
            asyncio.run(router.api_update_academic_final_material(
                "batch", body, Request({"type": "http", "headers": []}), {"id": 1},
            ))
        self.assertEqual(409, failure.exception.status_code)
        persist.assert_not_called()

    def test_regeneration_uses_snapshot_guard_and_returns_fresh_editor_version(self) -> None:
        from classroom_app.routers.materials_parts import academic_final_materials as router

        record = {"id": 88, "updated_at": "before", "export_payload_json": json.dumps({"document_type": service.ACADEMIC_EXAM_ANALYSIS_TYPE, "fields": {}, "structured": {}})}
        latest = {"id": 88, "updated_at": "after", "fields": {}, "structured": {"analysis_text": "新分析"}}
        connection_context = MagicMock()
        with (
            patch.object(router, "get_db_connection", return_value=connection_context),
            patch.object(router, "_batch_for_teacher", return_value={"class_offering_id": 3, "grade_record_id": 87, "analysis_record_id": 88}),
            patch.object(router, "_record_for_teacher", return_value=record),
            patch.object(router, "_load_course_analysis_context", return_value={}),
            patch.object(router, "_ai_review_and_analysis", new_callable=AsyncMock, return_value=("新分析", [], True)),
            patch.object(router, "_persist_final_material_record_update", new_callable=AsyncMock, return_value={"id": 88}) as persist,
            patch.object(router, "_serialize_record_payload", return_value=latest),
        ):
            result = asyncio.run(router.api_regenerate_academic_final_analysis(
                "batch", router.AcademicFinalMaterialRegenerateRequest(expected_updated_at="before"), {"id": 1},
            ))
        self.assertTrue(persist.call_args.kwargs["require_unchanged_record"])
        self.assertEqual(latest, result["record"])


if __name__ == "__main__":
    unittest.main()
