"""Run inside the fixed Debian image, with synthetic content only."""
import json
from pathlib import Path
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from PIL import Image

root = Path('/workspace')
doc = Document()
doc.add_heading('Synthetic DSH artifact', 0)
doc.add_paragraph('No platform data or real user content.')
doc.save(root / 'fixture.docx')
assert Document(root / 'fixture.docx').paragraphs[1].text == 'No platform data or real user content.'
workbook = Workbook()
workbook.active.append(['Value', 42])
workbook.save(root / 'fixture.xlsx')
assert load_workbook(root / 'fixture.xlsx').active.cell(1, 2).value == 42
slides = Presentation()
slide = slides.slides.add_slide(slides.slide_layouts[1])
slide.shapes.title.text = 'Synthetic fixture slide'
slides.save(root / 'fixture.pptx')
assert Presentation(root / 'fixture.pptx').slides[0].shapes.title.text == 'Synthetic fixture slide'
pdf = canvas.Canvas(str(root / 'fixture.pdf'))
pdf.drawString(72, 760, 'Synthetic DSH artifact PDF')
pdf.save()
assert (root / 'fixture.pdf').read_bytes().startswith(b'%PDF-')
Image.new('RGB', (40, 40), 'white').save(root / 'fixture.png')
with Image.open(root / 'fixture.png') as picture:
    assert picture.size == (40, 40)
print(json.dumps({'docx': 'generated_and_reopened', 'xlsx': 'generated_and_reopened',
                  'pptx': 'generated_and_reopened', 'pdf': 'generated_header_checked_not_visual_QA',
                  'png': 'generated_and_reopened'}))
