"""Runner files can change concurrently; never parse or return replacement bytes."""
from pathlib import Path
import os
import tempfile
import unittest
from unittest.mock import patch

from fastapi import HTTPException

from classroom_app.services import agent_bridge_service as bridge
from classroom_app.services import agent_platform_multipart_service as multipart
from classroom_app.services import agent_scoped_read_service as scoped


class AgentFileSnapshotTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        guard = patch.object(bridge, "allowed_file_roots", return_value=[self.root])
        guard.start()
        self.addCleanup(guard.stop)

    def test_regular_file_and_size_bound_are_checked_against_actual_handle(self):
        source = self.root / "report.md"
        source.write_bytes(b"actual text")
        self.assertEqual("actual text", bridge.read_platform_file(str(source))["content"])
        source.write_bytes(b"x" * (bridge.MAX_FILE_BYTES + 1))
        with self.assertRaises(ValueError):
            bridge.read_platform_file(str(source))

    def test_document_parser_reads_private_snapshot_after_original_is_replaced(self):
        from docx import Document
        from ai_assistant_doc_extract import extract_document_text

        source = self.root / "source.docx"
        document = Document()
        document.add_paragraph("Original authorized content")
        document.save(source)
        snapshots = []

        def replace_then_extract(snapshot, *args, **kwargs):
            snapshots.append(snapshot)
            self.assertNotEqual(source, snapshot)
            replacement = Document()
            replacement.add_paragraph("Replacement must not escape")
            replacement.save(source)
            return extract_document_text(snapshot, *args, **kwargs)

        with patch("ai_assistant_doc_extract.extract_document_text", side_effect=replace_then_extract):
            result = bridge.read_platform_file(str(source))
        self.assertIn("Original authorized content", result["content"])
        self.assertNotIn("Replacement", result["content"])
        self.assertFalse(snapshots[0].exists())

    def test_source_mutation_during_read_is_rejected(self):
        source = self.root / "changing.md"
        source.write_text("before", encoding="utf-8")
        actual_fstat = os.fstat
        calls = 0

        def change_before_final_stat(descriptor):
            nonlocal calls
            calls += 1
            if calls == 2:
                source.write_text("changed length after original read", encoding="utf-8")
            return actual_fstat(descriptor)

        with patch.object(bridge.os, "fstat", side_effect=change_before_final_stat):
            with self.assertRaises(ValueError):
                bridge.read_platform_file(str(source))

    def test_task_directory_redirection_never_changes_owner_root(self):
        own = self.root / "tasks" / "10"
        foreign = self.root / "tasks" / "11"
        own.mkdir(parents=True)
        foreign.mkdir()
        (foreign / "report.md").write_text("Other user private content", encoding="utf-8")
        try:
            (own / "redirect").symlink_to(foreign, target_is_directory=True)
        except OSError as exc:
            self.skipTest(f"Host cannot create symlink: {exc}")
        with patch.object(scoped, "AGENT_TASK_WORKSPACE_ROOT", self.root):
            with self.assertRaises((ValueError, HTTPException)):
                scoped.read_scoped_file(None, None, 10, path="redirect/report.md")

    @unittest.skipUnless(hasattr(os, "mkfifo") and os.open in os.supports_dir_fd, "requires POSIX FIFO")
    def test_fifo_without_writer_is_rejected_without_blocking(self):
        import concurrent.futures

        source = self.root / "pipe.md"
        os.mkfifo(source)
        with concurrent.futures.ThreadPoolExecutor() as pool:
            future = pool.submit(bridge.read_platform_file, str(source))
            with self.assertRaises(ValueError):
                future.result(timeout=1)
