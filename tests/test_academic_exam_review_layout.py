from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from classroom_app.services.academic_final_material_document_service import build_exam_analysis_docx


class AcademicExamReviewLayoutTests(unittest.TestCase):
    def document(self, fields=None, text="一、成绩分析\n教学分析内容。"):
        return Document(BytesIO(build_exam_analysis_docx({
            "fields": fields or {}, "structured": {"analysis_text": text},
        })))

    def test_reviews_keep_native_rows_and_note_in_original_table(self):
        document = self.document({"department_review_opinion": "已核", "dean_review_opinion": "同意"})
        row = document.tables[0].rows[19]
        self.assertEqual(2, len(row._tr.tc_lst))
        self.assertEqual(["已核", "同意"], [node.get("descr") for node in document._element.xpath(".//wp:docPr")])
        for cell in [row.cells[0], row.cells[-1]]:
            self.assertIn("审核意见：", cell.paragraphs[0].text)
        for cell in document.tables[0].rows[21].cells:
            self.assertIn("签字：", cell.text)
        self.assertIn("本表一式两份", document.tables[0].rows[-1].cells[0].text)
        self.assertFalse(document.paragraphs[-1]._p.xpath(".//w:pBdr"))

    def test_differently_shaped_signatures_fit_without_stretching_before_label(self):
        with TemporaryDirectory() as directory:
            paths = [Path(directory) / name for name in ["wide.png", "tall.png"]]
            sizes = [(1000, 40), (40, 1000)]
            for path, size in zip(paths, sizes):
                Image.new("RGBA", size, "black").save(path)
            document = self.document(dict(zip(["department_signature_image_path", "dean_signature_image_path"], map(str, paths))))
        anchors = document._element.xpath(".//wp:anchor")
        self.assertEqual(2, len(anchors))
        for anchor, size in zip(anchors, sizes):
            extent = anchor.find(qn("wp:extent"))
            width, height = int(extent.get("cx")) / 12700, int(extent.get("cy")) / 12700
            self.assertAlmostEqual(size[0] / size[1], width / height, delta=(size[0] / size[1]) * .0001)
            self.assertLessEqual(width, 88.01)
            self.assertLessEqual(height, 30.01)
            self.assertIsNotNone(anchor.find(qn("wp:wrapNone")))
        self.assertEqual(93, document.tables[0].rows[20].height.pt)
        self.assertEqual(16, document.tables[0].rows[21].height.pt)
        for cell in document.tables[0].rows[21].cells:
            self.assertEqual("签字：        ", cell.text)

    def test_long_body_is_rejected_without_losing_saved_content(self):
        text = "一、成绩分析\n" + "教学分析内容完整保留。" * 250
        payload = {"structured": {"analysis_text": text}}
        with self.assertRaisesRegex(ValueError, "超过原版单元格"):
            build_exam_analysis_docx(payload)
        self.assertEqual(text, payload["structured"]["analysis_text"])

    def test_explicit_empty_opinion_suppresses_old_opinion_image(self):
        asset = Path(__file__).resolve().parents[1] / "classroom_app/services/assets/gxufl_exam_review_checked.png"
        document = self.document({"department_review_opinion": "", "department_review_opinion_image_path": str(asset)})
        self.assertFalse(document._element.xpath(".//wp:anchor"))


if __name__ == "__main__":
    unittest.main()
