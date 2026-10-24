"""Locator-aware text extraction for the material knowledge index.

Turns one ``course_materials`` file row into searchable chunks, each tagged
with *where it is* inside the document so a hit can be reported as
"《文档》第 3 页" and opened at that spot:

======================  ===========  ===================================================
Material kind           locator      deep link
======================  ===========  ===================================================
PDF                     page         ``/materials/document-preview/{id}?page=N``
Word (.docx)            page/section ``/materials/document-preview/{id}?page=N`` (page from
                                     Word's last rendered page breaks; section label otherwise)
PowerPoint (.pptx)      slide        ``/materials/view/{id}`` (slide number in the label)
Excel (.xlsx)           sheet        ``/materials/document-preview/{id}``
Markdown                section      ``/materials/view/{id}#<heading-slug>`` (same slug rule
                                     as the viewer's TOC)
HTML / HTML package     lesson       ``/materials/render-view/{root}?path=lesson_N/...``
text / code             part         ``/materials/view/{id}``
======================  ===========  ===================================================

Pure function of the row + file bytes (plus a read-only package lookup for
HTML packages). No AI, no network. Everything is bounded so a single huge
upload cannot stall the scheduler worker.
"""
from __future__ import annotations

import re
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable
from urllib.parse import quote

from .file_service import resolve_global_file_path
from .html_package_service import (
    _row_get,
    extract_html_text,
    find_html_package_root,
    lesson_number_from_entry_name,
    package_relative_path,
)
from .material_document_kind_service import classify_material_kind

EXTRACTOR_VERSION = 1

MAX_FILE_BYTES = 60 * 1024 * 1024
MAX_ZIP_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 400 * 1024 * 1024
MAX_PDF_PAGES = 600
MAX_SLIDES = 600
MAX_SHEET_ROWS = 400
MAX_TOTAL_CHARS = 600_000
CHUNK_TARGET_CHARS = 700
CHUNK_MAX_CHARS = 1200
MIN_CHUNK_CHARS = 8

STATUS_READY = "ready"
STATUS_EMPTY = "empty"
STATUS_UNSUPPORTED = "unsupported"
STATUS_FAILED = "failed"

_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "gbk")
_MD_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")
_HTML_HEADING_RE = re.compile(r"<h([1-3])\b[^>]*>(.*?)</h\1\s*>", re.IGNORECASE | re.DOTALL)
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f​‌‍﻿]")
_WS_RE = re.compile(r"[ \t　]+")

_NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_P = "http://schemas.openxmlformats.org/presentationml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"


@dataclass
class MaterialChunk:
    ordinal: int
    locator_kind: str
    locator_label: str
    locator_url: str
    text: str
    page_number: int = 0
    anchor: str = ""
    heading: str = ""


@dataclass
class MaterialExtraction:
    status: str
    kind: str
    locator_kind: str = ""
    chunks: list[MaterialChunk] = field(default_factory=list)
    page_count: int = 0
    error: str = ""

    @property
    def char_count(self) -> int:
        return sum(len(chunk.text) for chunk in self.chunks)


# --------------------------------------------------------------------------
# text helpers
# --------------------------------------------------------------------------
def normalize_text(value: Any) -> str:
    """Strip control characters, collapse horizontal whitespace, keep line breaks."""
    text = _CONTROL_RE.sub("", str(value or "")).replace("\r\n", "\n").replace("\r", "\n")
    lines = [_WS_RE.sub(" ", line).strip() for line in text.split("\n")]
    cleaned: list[str] = []
    blank = False
    for line in lines:
        if not line:
            if not blank and cleaned:
                cleaned.append("")
            blank = True
            continue
        blank = False
        cleaned.append(line)
    return "\n".join(cleaned).strip()


def decode_text_bytes(raw: bytes) -> str:
    if b"\x00" in raw[:4096]:
        if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
            try:
                return raw.decode("utf-16")
            except UnicodeDecodeError:
                return ""
        return ""
    for encoding in _TEXT_ENCODINGS:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def split_text_into_pieces(text: str, *, target: int = CHUNK_TARGET_CHARS, maximum: int = CHUNK_MAX_CHARS) -> list[str]:
    """Cut text at paragraph / sentence boundaries into pieces of roughly ``target`` chars."""
    text = normalize_text(text)
    if not text:
        return []
    if len(text) <= maximum:
        return [text]
    paragraphs = [part.strip() for part in re.split(r"\n{2,}|\n", text) if part.strip()]
    pieces: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        units = [paragraph] if len(paragraph) <= maximum else _split_long_paragraph(paragraph, maximum)
        for unit in units:
            if not buffer:
                buffer = unit
            elif len(buffer) + 1 + len(unit) <= target:
                buffer = f"{buffer}\n{unit}"
            else:
                pieces.append(buffer)
                buffer = unit
    if buffer:
        pieces.append(buffer)
    return pieces


def _split_long_paragraph(paragraph: str, maximum: int) -> list[str]:
    sentences = re.split(r"(?<=[。！？!?；;])\s*", paragraph)
    units: list[str] = []
    buffer = ""
    for sentence in sentences:
        if not sentence:
            continue
        while len(sentence) > maximum:
            units.append(sentence[:maximum])
            sentence = sentence[maximum:]
        if not buffer:
            buffer = sentence
        elif len(buffer) + len(sentence) <= maximum:
            buffer += sentence
        else:
            units.append(buffer)
            buffer = sentence
    if buffer:
        units.append(buffer)
    return units


def markdown_heading_slug(text: str) -> str:
    """Mirror ``slugify`` in ``static/js/material_viewer.js`` (ASCII-only word characters)."""
    slug = re.sub(r"[^a-z0-9_]+", "-", str(text or "").lower()).strip("-")
    return slug or "section"


class _SlugCounter:
    def __init__(self) -> None:
        self._seen: dict[str, int] = {}

    def next(self, heading: str) -> str:
        base = markdown_heading_slug(heading)
        count = self._seen.get(base, 0)
        self._seen[base] = count + 1
        return base if count == 0 else f"{base}-{count + 1}"


def with_query(url: str, **params: Any) -> str:
    """Append query parameters to ``url`` while keeping any ``#fragment`` last."""
    extras = [f"{key}={quote(str(value), safe='/')}" for key, value in params.items() if value not in (None, "", 0)]
    if not extras:
        return url
    base, _, fragment = url.partition("#")
    joiner = "&" if "?" in base else "?"
    rebuilt = f"{base}{joiner}{'&'.join(extras)}"
    return f"{rebuilt}#{fragment}" if fragment else rebuilt


# --------------------------------------------------------------------------
# section → chunk assembly
# --------------------------------------------------------------------------
@dataclass
class _Section:
    text: str
    locator_kind: str
    locator_label: str
    locator_url: str
    page_number: int = 0
    anchor: str = ""
    heading: str = ""


def _assemble(sections: Iterable[_Section]) -> list[MaterialChunk]:
    chunks: list[MaterialChunk] = []
    total = 0
    for section in sections:
        pieces = split_text_into_pieces(section.text)
        for piece in pieces:
            if len(piece) < MIN_CHUNK_CHARS:
                continue
            if total + len(piece) > MAX_TOTAL_CHARS:
                return chunks
            total += len(piece)
            chunks.append(
                MaterialChunk(
                    ordinal=len(chunks),
                    locator_kind=section.locator_kind,
                    locator_label=section.locator_label,
                    locator_url=section.locator_url,
                    text=piece,
                    page_number=int(section.page_number or 0),
                    anchor=section.anchor,
                    heading=section.heading[:200],
                )
            )
    return chunks


def _heading_sections(
    lines: Iterable[tuple[str | None, str]],
    *,
    locator_kind: str,
    url_for: Callable[[str, int], str],
    label_for: Callable[[str, int], str],
    anchor_for: Callable[[str], str] | None = None,
) -> list[_Section]:
    """Group ``(heading_or_None, text)`` pairs into sections; ``None`` heading = preamble."""
    sections: list[_Section] = []
    index = 0
    for heading, text in lines:
        body = normalize_text(text)
        heading_text = normalize_text(heading or "").replace("\n", " ")
        if not body and not heading_text:
            continue
        index += 1
        anchor = anchor_for(heading_text) if (anchor_for and heading_text) else ""
        full_text = f"{heading_text}\n{body}".strip() if heading_text else body
        sections.append(
            _Section(
                text=full_text,
                locator_kind=locator_kind,
                locator_label=label_for(heading_text, index),
                locator_url=url_for(anchor, index),
                anchor=anchor,
                heading=heading_text,
            )
        )
    return sections


# --------------------------------------------------------------------------
# per-kind extractors
# --------------------------------------------------------------------------
def _split_markdown(text: str) -> list[tuple[str | None, str]]:
    pairs: list[tuple[str | None, str]] = []
    heading: str | None = None
    buffer: list[str] = []
    in_fence = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_fence = not in_fence
            buffer.append(line)
            continue
        match = None if in_fence else _MD_HEADING_RE.match(stripped)
        if match:
            if heading is not None or any(part.strip() for part in buffer):
                pairs.append((heading, "\n".join(buffer)))
            heading = match.group(2).strip()
            buffer = []
            continue
        buffer.append(line)
    if heading is not None or any(part.strip() for part in buffer):
        pairs.append((heading, "\n".join(buffer)))
    return pairs


def _extract_markdown(material_id: int, text: str) -> MaterialExtraction:
    counter = _SlugCounter()
    pairs = _split_markdown(text)
    if not pairs:
        return MaterialExtraction(STATUS_EMPTY, "markdown")
    view_url = f"/materials/view/{material_id}"
    sections = _heading_sections(
        pairs,
        locator_kind="section",
        url_for=lambda anchor, _index: f"{view_url}#{anchor}" if anchor else view_url,
        label_for=lambda heading, index: f"章节「{heading[:60]}」" if heading else "开头部分",
        anchor_for=counter.next,
    )
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, "markdown", "section", chunks, 0)


def _extract_plain_text(material_id: int, text: str, kind: str) -> MaterialExtraction:
    view_url = f"/materials/view/{material_id}"
    pieces = split_text_into_pieces(text)
    sections = [
        _Section(text=piece, locator_kind="part", locator_label=f"第 {index} 段", locator_url=view_url)
        for index, piece in enumerate(pieces, start=1)
    ]
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, kind, "part", chunks, 0)


def _split_html_sections(raw_html: str) -> list[tuple[str | None, str]]:
    pairs: list[tuple[str | None, str]] = []
    cursor = 0
    heading: str | None = None
    for match in _HTML_HEADING_RE.finditer(raw_html):
        segment = raw_html[cursor : match.start()]
        text = extract_html_text(segment, max_chars=MAX_TOTAL_CHARS)
        if text or heading is not None:
            pairs.append((heading, text))
        heading = extract_html_text(match.group(2), max_chars=400).replace("\n", " ").strip() or None
        cursor = match.end()
    tail = extract_html_text(raw_html[cursor:], max_chars=MAX_TOTAL_CHARS)
    if tail or heading is not None:
        pairs.append((heading, tail))
    return pairs


def _extract_html(conn: Any, material: dict[str, Any], raw_html: str) -> MaterialExtraction:
    material_id = int(material["id"])
    package = None
    try:
        package = find_html_package_root(conn, material) if conn is not None else None
    except Exception:  # noqa: BLE001 — package detection is a bonus, never a blocker
        package = None

    lesson_number = 0
    prefix = ""
    base_url = f"/materials/render-view/{material_id}"
    locator_kind = "section"
    if package:
        relpath = package_relative_path(package, material.get("material_path"))
        base_url = with_query(f"/materials/render-view/{int(package['root_node_id'])}", path=relpath)
        main_entry = package.get("main_entry") or {}
        if int(_row_get(main_entry, "id") or 0) == material_id:
            prefix = "课程首页"
            locator_kind = "lesson"
        else:
            lesson_number = lesson_number_from_entry_name(material.get("name"))
            if lesson_number <= 0:
                for lesson in package.get("lessons") or []:
                    if int(_row_get(lesson.get("entry") or {}, "id") or 0) == material_id:
                        lesson_number = int(lesson["number"])
                        break
            if lesson_number > 0:
                prefix = f"第 {lesson_number} 次课"
                locator_kind = "lesson"

    def label_for(heading: str, index: int) -> str:
        head = f"「{heading[:60]}」" if heading else (f"第 {index} 部分" if index > 1 else "开头部分")
        return f"{prefix} · {head}" if prefix else head

    if 'data-lessondoc="' in raw_html[:2000]:
        pairs: list[tuple[str | None, str]] = [(None, extract_html_text(raw_html, max_chars=MAX_TOTAL_CHARS))]
    else:
        pairs = _split_html_sections(raw_html)
    sections = _heading_sections(
        pairs,
        locator_kind=locator_kind,
        url_for=lambda _anchor, _index: base_url,
        label_for=label_for,
    )
    for section in sections:
        section.page_number = lesson_number
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, "html", locator_kind, chunks, lesson_number)


def _extract_pdf(material_id: int, path: Path) -> MaterialExtraction:
    import fitz  # PyMuPDF (in requirements.lock.txt)

    sections: list[_Section] = []
    with fitz.open(str(path)) as document:
        page_count = int(len(document))
        for index in range(min(page_count, MAX_PDF_PAGES)):
            page_number = index + 1
            try:
                text = document[index].get_text("text")
            except Exception:  # noqa: BLE001 — one broken page must not sink the file
                continue
            text = normalize_text(text)
            if not text:
                continue
            sections.append(
                _Section(
                    text=text,
                    locator_kind="page",
                    locator_label=f"第 {page_number} 页",
                    locator_url=f"/materials/document-preview/{material_id}?page={page_number}",
                    page_number=page_number,
                )
            )
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, "pdf", "page", chunks, page_count)


def _docx_paragraph_page_breaks(paragraph) -> int:
    count = 0
    for element in paragraph._p.iter():  # noqa: SLF001 — python-docx exposes no public page-break API
        tag = element.tag
        if tag == f"{{{_NS_W}}}lastRenderedPageBreak":
            count += 1
        elif tag == f"{{{_NS_W}}}br" and element.get(f"{{{_NS_W}}}type") == "page":
            count += 1
    return count


def _extract_docx(material_id: int, path: Path) -> MaterialExtraction:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    state = {"page": 1, "saw_break": False, "heading": "", "buffer": [], "index": 0}
    sections: list[_Section] = []

    def flush() -> None:
        body = normalize_text("\n".join(state["buffer"]))
        heading = state["heading"]
        state["buffer"] = []
        if not body and not heading:
            return
        state["index"] += 1
        page = state["page"]
        paged = state["saw_break"]
        label = f"第 {page} 页" if paged else (f"「{heading[:60]}」" if heading else f"第 {state['index']} 节")
        url = f"/materials/document-preview/{material_id}"
        if paged:
            url = f"{url}?page={page}"
        sections.append(
            _Section(
                text=f"{heading}\n{body}".strip() if heading else body,
                locator_kind="page" if paged else "section",
                locator_label=label,
                locator_url=url,
                page_number=page if paged else 0,
                heading=heading,
            )
        )

    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            breaks = _docx_paragraph_page_breaks(paragraph)
            style_name = str(getattr(getattr(paragraph, "style", None), "name", "") or "")
            lowered_style = style_name.lower()
            is_heading = lowered_style.startswith("heading") or style_name.startswith("标题") or lowered_style == "title"
            text = paragraph.text or ""
            if breaks:
                state["saw_break"] = True
                flush()
                state["heading"] = "" if not is_heading else state["heading"]
                state["page"] += breaks
            if is_heading and text.strip():
                flush()
                state["heading"] = text.strip()
                continue
            if text.strip():
                state["buffer"].append(text)
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row in table.rows:
                cells = [normalize_text(cell.text).replace("\n", " ") for cell in row.cells]
                line = " | ".join(cell for cell in cells if cell)
                if line:
                    state["buffer"].append(line)
    flush()
    chunks = _assemble(sections)
    paged = bool(state["saw_break"])
    return MaterialExtraction(
        STATUS_READY if chunks else STATUS_EMPTY,
        "docx",
        "page" if paged else "section",
        chunks,
        int(state["page"]) if paged else 0,
    )


def _ordered_slide_names(archive: zipfile.ZipFile) -> list[str]:
    names = archive.namelist()
    slide_names = [name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
    try:
        rels_root = ET.fromstring(archive.read("ppt/_rels/presentation.xml.rels"))
        rel_targets = {
            rel.get("Id"): rel.get("Target") for rel in rels_root.iter(f"{{{_NS_REL}}}Relationship")
        }
        presentation = ET.fromstring(archive.read("ppt/presentation.xml"))
        ordered: list[str] = []
        for slide_id in presentation.iter(f"{{{_NS_P}}}sldId"):
            target = rel_targets.get(slide_id.get(f"{{{_NS_R}}}id") or "")
            if not target:
                continue
            normalized = target.lstrip("/")
            if not normalized.startswith("ppt/"):
                normalized = f"ppt/{normalized}"
            if normalized in names:
                ordered.append(normalized)
        if ordered:
            return ordered
    except (KeyError, ET.ParseError):
        pass
    return sorted(slide_names, key=lambda name: int(re.search(r"slide(\d+)", name).group(1)))


def _extract_pptx(material_id: int, path: Path) -> MaterialExtraction:
    sections: list[_Section] = []
    with zipfile.ZipFile(path, "r") as archive:
        slide_names = _ordered_slide_names(archive)
        for index, name in enumerate(slide_names[:MAX_SLIDES], start=1):
            try:
                root = ET.fromstring(archive.read(name))
            except (KeyError, ET.ParseError):
                continue
            lines = [node.text for node in root.iter(f"{{{_NS_A}}}t") if node.text and node.text.strip()]
            text = normalize_text("\n".join(lines))
            if not text:
                continue
            sections.append(
                _Section(
                    text=text,
                    locator_kind="slide",
                    locator_label=f"第 {index} 页（幻灯片）",
                    locator_url=f"/materials/view/{material_id}",
                    page_number=index,
                    heading=lines[0].strip()[:120] if lines else "",
                )
            )
        slide_total = len(slide_names)
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, "pptx", "slide", chunks, slide_total)


def _extract_xlsx(material_id: int, path: Path) -> MaterialExtraction:
    import openpyxl

    sections: list[_Section] = []
    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
            lines: list[str] = []
            for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if row_index > MAX_SHEET_ROWS:
                    break
                cells = [normalize_text(cell).replace("\n", " ") for cell in row if cell not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
            text = normalize_text("\n".join(lines))
            if not text:
                continue
            sections.append(
                _Section(
                    text=text,
                    locator_kind="sheet",
                    locator_label=f"工作表「{str(sheet.title)[:40]}」",
                    locator_url=f"/materials/document-preview/{material_id}",
                    page_number=sheet_index,
                    heading=str(sheet.title or ""),
                )
            )
        sheet_total = len(workbook.worksheets)
    finally:
        workbook.close()
    chunks = _assemble(sections)
    return MaterialExtraction(STATUS_READY if chunks else STATUS_EMPTY, "xlsx", "sheet", chunks, sheet_total)


# --------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------
def extract_material_chunks(conn: Any, material_row: Any) -> MaterialExtraction:
    """Extract locator-tagged chunks for one ``course_materials`` file row.

    Never raises: unsupported / unreadable inputs come back as a status so the
    indexer can record the outcome and stop retrying until the file changes.
    """
    material = dict(material_row)
    material_id = int(material.get("id") or 0)
    kind = classify_material_kind(material)
    if str(material.get("node_type") or "") != "file":
        return MaterialExtraction(STATUS_UNSUPPORTED, kind, error="folder")
    ext = str(material.get("file_ext") or "").strip().lower().lstrip(".")
    if not ext and "." in str(material.get("name") or ""):
        ext = str(material["name"]).rsplit(".", 1)[-1].lower()

    try:
        size = int(material.get("file_size") or 0)
    except (TypeError, ValueError):
        size = 0
    if size > MAX_FILE_BYTES:
        return MaterialExtraction(STATUS_UNSUPPORTED, kind, error="too_large")

    if kind in {"image", "media", "archive", "folder", "git_repository"}:
        return MaterialExtraction(STATUS_UNSUPPORTED, kind, error=f"kind:{kind}")
    if ext in {"doc", "ppt", "xls", "odt", "odp", "ods", "rtf"}:
        return MaterialExtraction(STATUS_UNSUPPORTED, kind, error=f"legacy:{ext}")

    path = resolve_global_file_path(str(material.get("file_hash") or "")) if material.get("file_hash") else None
    if path and ext in {"docx", "pptx", "xlsx"} and not _zip_members_within_budget(path):
        return MaterialExtraction(STATUS_UNSUPPORTED, kind, error="zip_member_too_large")
    try:
        if kind == "markdown" or ext in {"md", "markdown"}:
            text = _read_text(path) if path else ""
            if not text:
                text = str(material.get("ai_optimized_markdown") or "")
            return _extract_markdown(material_id, text)
        if kind in {"html_page", "html_package", "lessondoc_pack"} or ext in {"html", "htm", "xhtml"}:
            if not path:
                return MaterialExtraction(STATUS_UNSUPPORTED, kind, error="missing_file")
            return _extract_html(conn, material, _read_text(path))
        if not path:
            return MaterialExtraction(STATUS_UNSUPPORTED, kind, error="missing_file")
        if kind == "pdf" or ext == "pdf":
            return _extract_pdf(material_id, path)
        if ext == "docx":
            return _extract_docx(material_id, path)
        if ext == "pptx":
            return _extract_pptx(material_id, path)
        if ext == "xlsx":
            return _extract_xlsx(material_id, path)
        if kind in {"text", "code"} or ext in {"txt", "log", "csv", "json"}:
            text = _read_text(path)
            if not text:
                return MaterialExtraction(STATUS_EMPTY, kind)
            return _extract_plain_text(material_id, text, kind)
    except Exception as exc:  # noqa: BLE001 — recorded on the document row, never propagated
        return MaterialExtraction(STATUS_FAILED, kind, error=f"{type(exc).__name__}: {str(exc)[:200]}")
    return MaterialExtraction(STATUS_UNSUPPORTED, kind, error=f"ext:{ext or 'none'}")


def _zip_members_within_budget(path: Path) -> bool:
    """Reject Office zips whose *uncompressed* members would blow memory (zip bombs, huge XML)."""
    try:
        with zipfile.ZipFile(path, "r") as archive:
            total = 0
            for info in archive.infolist():
                if info.file_size > MAX_ZIP_MEMBER_BYTES:
                    return False
                total += info.file_size
                if total > MAX_ZIP_TOTAL_BYTES:
                    return False
    except (zipfile.BadZipFile, OSError):
        return False
    return True


def _read_text(path: Path | None) -> str:
    if not path:
        return ""
    try:
        return decode_text_bytes(path.read_bytes())
    except OSError:
        return ""
