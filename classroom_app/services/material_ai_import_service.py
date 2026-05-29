import asyncio
import base64
import csv
import json
import mimetypes
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Awaitable, Callable

from fastapi import HTTPException

from .file_preview_service import TEXT_CONTENT_ENCODINGS
from .material_final_document_service import (
    FINAL_MATERIAL_TYPES,
    extract_markdown_tables,
    normalize_final_material_payload,
)

try:
    from ai_assistant_doc_extract import extract_document_text, render_pdf_pages_to_data_urls
except Exception:
    extract_document_text = None
    render_pdf_pages_to_data_urls = None


MAX_EXTRACT_TEXT_BYTES = 900_000
MAX_AI_TEXT_CHARS = 120_000
MAX_VISION_IMAGES = 8
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MIN_USABLE_TEXT_CHARS = 40

MaterialAiChat = Callable[..., Awaitable[Any]]


MATERIAL_AI_IMPORT_GROUPS: list[dict[str, Any]] = [
    {
        "key": "teaching_material",
        "label": "教学材料",
        "description": "日常教学过程材料",
        "types": [
            {
                "key": "teaching_document",
                "label": "教学文档",
                "template_key": "teaching_document",
                "aliases": ["教学文档", "课程资料", "授课资料"],
            },
            {
                "key": "teaching_summary",
                "label": "工作总结",
                "template_key": "teaching_summary",
                "aliases": ["教师教学工作总结", "教学工作总结", "工作总结"],
            },
            {
                "key": "lesson_plan",
                "label": "教案",
                "template_key": "lesson_plan",
                "aliases": ["教案", "授课教案", "课程教案"],
            },
            {
                "key": "teaching_calendar",
                "label": "教学日历",
                "template_key": "teaching_calendar",
                "aliases": ["教学日历", "授课教师名单", "教学进度表"],
            },
            {
                "key": "evaluation_sheet",
                "label": "评学表",
                "template_key": "evaluation_sheet",
                "aliases": ["评学表", "课程评学表"],
            },
        ],
    },
    {
        "key": "final_material",
        "label": "期末材料",
        "description": "课程期末考核归档材料",
        "types": [
            {
                "key": "syllabus",
                "label": "教学大纲",
                "template_key": "final_syllabus",
                "aliases": ["教学大纲", "课程教学大纲", "大纲"],
            },
            {
                "key": "assessment_plan",
                "label": "考核计划表",
                "template_key": "assessment_plan",
                "aliases": ["考核计划表", "课程考核计划表", "非笔试考核计划表"],
            },
            {
                "key": "grading_rubric",
                "label": "评分细则",
                "template_key": "grading_rubric",
                "aliases": ["评分细则", "课程考核评分细则", "评分标准"],
            },
            {
                "key": "exam_paper",
                "label": "考核试卷",
                "template_key": "exam_paper",
                "aliases": ["考核试卷", "试卷", "课程考核试卷"],
            },
            {
                "key": "final_teaching_summary",
                "label": "教学工作总结",
                "template_key": "final_teaching_summary",
                "aliases": ["教学工作总结", "教师教学工作总结", "工作总结"],
            },
        ],
    },
]

_TYPE_INDEX: dict[tuple[str, str], dict[str, Any]] = {}
for _group in MATERIAL_AI_IMPORT_GROUPS:
    for _doc_type in _group["types"]:
        _TYPE_INDEX[(_group["key"], _doc_type["key"])] = {
            **_doc_type,
            "group_key": _group["key"],
            "group_label": _group["label"],
        }


@dataclass
class MaterialExtraction:
    text: str = ""
    method: str = ""
    source_kind: str = ""
    warnings: list[str] = field(default_factory=list)
    images: list[dict[str, str]] = field(default_factory=list)
    truncated: bool = False
    quality: dict[str, Any] = field(default_factory=dict)


@dataclass
class MaterialParseResult:
    metadata: dict[str, Any]
    content_markdown: str
    tables: list[dict[str, Any]]
    warnings: list[str]
    export_payload: dict[str, Any]
    raw_ai_result: dict[str, Any]
    parsed_payload: dict[str, Any]
    content_quality: dict[str, Any]
    extraction_method: str
    document_group: str
    document_type: str
    document_type_label: str
    ai_used: bool


def get_material_ai_import_registry() -> list[dict[str, Any]]:
    return [
        {
            "key": group["key"],
            "label": group["label"],
            "description": group.get("description", ""),
            "types": [
                {
                    "key": doc_type["key"],
                    "label": doc_type["label"],
                    "template_key": doc_type.get("template_key", ""),
                    "aliases": list(doc_type.get("aliases", [])),
                }
                for doc_type in group["types"]
            ],
        }
        for group in MATERIAL_AI_IMPORT_GROUPS
    ]


def resolve_material_ai_import_type(document_group: str, document_type: str) -> dict[str, Any]:
    group_key = str(document_group or "").strip()
    type_key = str(document_type or "").strip()
    type_meta = _TYPE_INDEX.get((group_key, type_key))
    if not type_meta:
        raise HTTPException(400, "材料解析类型不受支持")
    return type_meta.copy()


async def parse_material_document(
    *,
    file_path: Path,
    original_name: str,
    document_group: str,
    document_type: str,
    ai_chat: MaterialAiChat,
) -> MaterialParseResult:
    type_meta = resolve_material_ai_import_type(document_group, document_type)
    extraction = await asyncio.to_thread(extract_material_content, file_path, original_name)
    extraction.quality = _assess_text_quality(extraction.text, method=extraction.method)
    warnings = list(extraction.warnings)
    ai_used = False
    raw_ai_result: dict[str, Any] | None = None

    if extraction.truncated:
        warnings.append("本地抽取内容较长，已截断后交给 AI 解析。")
    if extraction.text.strip() and not extraction.quality.get("usable", False):
        warnings.extend(str(reason) for reason in extraction.quality.get("reasons", []) if reason)

    text_for_ai = _limit_text_for_ai(extraction.text) if extraction.quality.get("usable", False) else ""
    if text_for_ai.strip():
        try:
            raw_ai_result = await ai_chat(
                *_build_material_prompts(
                    original_name=original_name,
                    type_meta=type_meta,
                    extraction=extraction,
                    text_for_ai=text_for_ai,
                    vision_mode=False,
                ),
                capability="thinking",
                response_format="json",
                file_texts=[{"name": original_name, "content": text_for_ai}],
                task_type="deep_text_reasoning",
                task_priority="background",
                task_label="material_ai_import:text",
                timeout=240.0,
            )
            ai_used = True
        except Exception as exc:
            warnings.append(f"文本 AI 解析失败，已尝试兼容兜底: {_format_exception(exc)}")

    if not raw_ai_result and _needs_vision_fallback(extraction):
        image_inputs = extraction.images[:MAX_VISION_IMAGES]
        if image_inputs:
            try:
                raw_ai_result = await ai_chat(
                    *_build_material_prompts(
                        original_name=original_name,
                        type_meta=type_meta,
                        extraction=extraction,
                        text_for_ai=text_for_ai,
                        vision_mode=True,
                    ),
                    capability="vision",
                    response_format="json",
                    base64_urls=[item["data_url"] for item in image_inputs if item.get("data_url")],
                    task_type="deep_multimodal_reasoning",
                    task_priority="background",
                    task_label="material_ai_import:vision",
                    timeout=300.0,
                )
                ai_used = True
            except Exception as exc:
                warnings.append(f"视觉 AI 兜底失败: {_format_exception(exc)}")

    if not raw_ai_result:
        if extraction.method == "unsupported":
            raise HTTPException(415, "当前文件格式暂不支持自动解析，请先转换为 docx、xlsx 或 PDF 后重试。")
        if not extraction.text.strip() or not extraction.quality.get("usable", False):
            raise HTTPException(422, "无法从该材料中抽取可解析内容，请更换文件或先转为 PDF/Word/Excel 后重试。")
        raw_ai_result = _build_local_fallback_result(
            original_name=original_name,
            type_meta=type_meta,
            extraction=extraction,
            warning="AI 未返回有效解析结果，系统已保留本地抽取内容。",
        )
        ai_used = False

    result = normalize_ai_parse_result(
        raw_ai_result,
        original_name=original_name,
        type_meta=type_meta,
        extraction=extraction,
        extra_warnings=warnings,
        ai_used=ai_used,
    )
    if not result.content_quality.get("usable", False):
        raise HTTPException(422, "解析结果质量校验未通过，系统已阻止保存乱码内容。请将文件另存为 docx/PDF 后重试，或稍后使用视觉解析兜底。")
    return result


def extract_material_content(file_path: Path, original_name: str) -> MaterialExtraction:
    ext = Path(original_name or file_path.name).suffix.lower()
    if ext in {".md", ".markdown", ".txt", ".csv", ".json"}:
        return _extract_text_like(file_path, ext)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext == ".doc":
        return _extract_legacy_document(file_path, ext)
    if ext == ".xlsx":
        return _extract_xlsx(file_path)
    if ext == ".xls":
        return _extract_xls(file_path)
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}:
        return _extract_image_file(file_path, ext)

    if extract_document_text and ext in {".pptx", ".ppt"}:
        extracted = extract_document_text(file_path, ext, max_bytes=MAX_EXTRACT_TEXT_BYTES)
        return MaterialExtraction(
            text=extracted.text,
            method=f"{ext.lstrip('.')}_document_extract",
            source_kind=ext.lstrip("."),
            warnings=["该格式按通用文档抽取处理，复杂版式可能需要人工复核。"],
            images=list(extracted.images or [])[:MAX_VISION_IMAGES],
            truncated=bool(extracted.truncated),
        )

    return MaterialExtraction(
        text="",
        method="unsupported",
        source_kind=ext.lstrip(".") or "unknown",
        warnings=["当前文件格式暂未提供本地结构化抽取。"],
    )


def normalize_ai_parse_result(
    raw_result: Any,
    *,
    original_name: str,
    type_meta: dict[str, Any],
    extraction: MaterialExtraction,
    extra_warnings: list[str],
    ai_used: bool,
) -> MaterialParseResult:
    parsed = _coerce_json_object(raw_result)
    metadata = _coerce_dict(parsed.get("metadata"))
    content_markdown = str(parsed.get("content_markdown") or parsed.get("content") or "").strip()
    if not content_markdown:
        content_markdown = extraction.text.strip() if extraction.quality.get("usable", False) else ""
    if not content_markdown:
        raise HTTPException(422, "AI 未返回有效正文，系统已取消保存该解析结果。")

    metadata.setdefault("source_filename", original_name)
    metadata.setdefault("document_group", type_meta["group_label"])
    metadata.setdefault("document_type", type_meta["label"])
    _normalize_academic_period(metadata)

    tables = _coerce_tables(parsed.get("tables"))
    if not tables:
        tables = extract_markdown_tables(content_markdown)
    warnings = _merge_warnings(extra_warnings, parsed.get("warnings"))
    export_payload = _coerce_dict(parsed.get("export_payload"))
    if not export_payload:
        export_payload = build_material_export_payload(
            document_group=type_meta["group_key"],
            document_type=type_meta["key"],
            type_meta=type_meta,
            metadata=metadata,
            content_markdown=content_markdown,
            tables=tables,
        )
    else:
        export_payload.setdefault("document_group", type_meta["group_key"])
        export_payload.setdefault("document_type", type_meta["key"])
        export_payload.setdefault("document_type_label", type_meta["label"])
        export_payload.setdefault("template_key", type_meta.get("template_key", type_meta["key"]))
    if type_meta["key"] in FINAL_MATERIAL_TYPES:
        export_payload = normalize_final_material_payload(
            document_type=type_meta["key"],
            metadata=metadata,
            content_markdown=content_markdown,
            tables=tables,
            export_payload=export_payload,
        )
        metadata.update(
            {
                key: value
                for key, value in _coerce_dict(export_payload.get("fields")).items()
                if key not in metadata and value not in (None, "", [], {})
            }
        )

    content_quality = _assess_text_quality(content_markdown, method="parsed_markdown")
    raw_result_dict = parsed if isinstance(parsed, dict) else {}
    parsed_payload = {
        "metadata": metadata,
        "content_markdown": content_markdown,
        "tables": tables,
        "warnings": warnings,
        "export_payload": export_payload,
        "document_group": type_meta["group_key"],
        "document_type": type_meta["key"],
        "document_type_label": type_meta["label"],
        "extraction": {
            "method": extraction.method or "unknown",
            "source_kind": extraction.source_kind,
            "truncated": extraction.truncated,
            "quality": extraction.quality,
        },
        "content_quality": content_quality,
        "ai_used": ai_used,
    }
    return MaterialParseResult(
        metadata=metadata,
        content_markdown=content_markdown,
        tables=tables,
        warnings=warnings,
        export_payload=export_payload,
        raw_ai_result=raw_result_dict,
        parsed_payload=parsed_payload,
        content_quality=content_quality,
        extraction_method=extraction.method or "unknown",
        document_group=type_meta["group_key"],
        document_type=type_meta["key"],
        document_type_label=type_meta["label"],
        ai_used=ai_used,
    )


def build_import_readme(
    *,
    result: MaterialParseResult,
    original_name: str,
) -> str:
    title = _metadata_title(result.metadata, original_name, result.document_type_label)
    content = _strip_duplicate_title(result.content_markdown.strip(), title)
    if re.match(r"^#\s+", content):
        return content.strip() + "\n"
    return f"# {title}\n\n{content}".strip() + "\n"


def build_material_export_payload(
    *,
    document_group: str,
    document_type: str,
    type_meta: dict[str, Any],
    metadata: dict[str, Any],
    content_markdown: str,
    tables: list[dict[str, Any]],
) -> dict[str, Any]:
    payload = {
        "document_group": document_group,
        "document_type": document_type,
        "document_type_label": type_meta.get("label", document_type),
        "template_key": type_meta.get("template_key", document_type),
        "fields": metadata,
        "sections": _infer_export_sections(document_type, content_markdown),
        "tables": tables,
        "compatibility": {
            "source_format_preserved": False,
            "requires_template_confirmation": True,
            "layout_source": "parsed_content",
        },
    }
    if document_type in FINAL_MATERIAL_TYPES:
        payload = normalize_final_material_payload(
            document_type=document_type,
            metadata=metadata,
            content_markdown=content_markdown,
            tables=tables,
            export_payload=payload,
        )
    return payload


def _extract_text_like(file_path: Path, ext: str) -> MaterialExtraction:
    raw = file_path.read_bytes()
    for encoding in TEXT_CONTENT_ENCODINGS:
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="ignore")

    if ext == ".csv":
        text = _csv_to_markdown(text)

    truncated_text, truncated = _truncate_by_bytes(text, MAX_EXTRACT_TEXT_BYTES)
    return MaterialExtraction(
        text=truncated_text,
        method="plain_text_decode",
        source_kind=ext.lstrip("."),
        truncated=truncated,
    )


def _extract_docx(file_path: Path) -> MaterialExtraction:
    warnings: list[str] = []
    parts: list[str] = []
    table_index = 0
    try:
        from docx import Document
        from docx.oxml.table import CT_Tbl
        from docx.oxml.text.paragraph import CT_P
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = Document(str(file_path))
        for child in document.element.body.iterchildren():
            if isinstance(child, CT_P):
                text = Paragraph(child, document).text.strip()
                if text:
                    parts.append(text)
            elif isinstance(child, CT_Tbl):
                table = Table(child, document)
                table_index += 1
                table_markdown = _docx_table_to_markdown(table)
                if table_markdown:
                    parts.append(f"\n[表格 {table_index}]\n{table_markdown}")
    except Exception as exc:
        warnings.append(f"python-docx 结构化抽取失败: {_format_exception(exc)}")
        fallback = _extract_docx_zip_text(file_path)
        if fallback:
            parts.append(fallback)

    images: list[dict[str, str]] = []
    truncated = False
    if extract_document_text:
        extracted = extract_document_text(file_path, ".docx", max_bytes=MAX_EXTRACT_TEXT_BYTES)
        images = list(extracted.images or [])[:MAX_VISION_IMAGES]
        if not parts and extracted.text:
            parts.append(extracted.text)
        truncated = bool(extracted.truncated)

    text, local_truncated = _truncate_by_bytes("\n\n".join(parts), MAX_EXTRACT_TEXT_BYTES)
    if len(text.strip()) < 800 and not images:
        images.extend(_render_office_pages_to_images(file_path, ".docx", warnings))
    return MaterialExtraction(
        text=text,
        method="python_docx_tables",
        source_kind="docx",
        warnings=warnings,
        images=images,
        truncated=truncated or local_truncated,
    )


def _extract_legacy_document(file_path: Path, ext: str) -> MaterialExtraction:
    warnings = ["旧版 Word 文档按兼容模式抽取，系统会优先使用 antiword/LibreOffice，避免保存二进制乱码。"]
    antiword_text = _extract_doc_with_antiword(file_path, warnings)
    if _assess_text_quality(antiword_text, method="antiword_doc_extract").get("usable", False):
        text, truncated = _truncate_by_bytes(antiword_text, MAX_EXTRACT_TEXT_BYTES)
        return MaterialExtraction(
            text=text,
            method="antiword_doc_extract",
            source_kind=ext.lstrip("."),
            warnings=warnings,
            truncated=truncated,
        )

    converted = _extract_legacy_doc_via_libreoffice(file_path, warnings)
    if converted.text.strip() or converted.images:
        converted.warnings = _merge_warnings(warnings, converted.warnings)
        converted.source_kind = ext.lstrip(".")
        return converted

    if extract_document_text:
        extracted = extract_document_text(file_path, ext, max_bytes=MAX_EXTRACT_TEXT_BYTES)
        images = list(extracted.images or [])[:MAX_VISION_IMAGES]
        if len(str(extracted.text or "").strip()) < 800 and not images:
            images.extend(_render_office_pages_to_images(file_path, ext, warnings))
        binary_quality = _assess_text_quality(extracted.text, method="legacy_doc_binary_extract")
        text = extracted.text if binary_quality.get("usable", False) else ""
        if extracted.text and not text:
            warnings.extend(binary_quality.get("reasons", []))
            warnings.append("已丢弃旧版 Word 二进制兜底文本，防止乱码进入材料库。")
        return MaterialExtraction(
            text=text,
            method="legacy_doc_binary_extract" if text else "legacy_doc_unreadable",
            source_kind=ext.lstrip("."),
            warnings=warnings,
            images=images[:MAX_VISION_IMAGES],
            truncated=bool(extracted.truncated),
        )
    return MaterialExtraction(method="legacy_doc_unavailable", source_kind=ext.lstrip("."), warnings=warnings)


def _extract_doc_with_antiword(file_path: Path, warnings: list[str]) -> str:
    antiword = shutil.which("antiword")
    if not antiword:
        warnings.append("未检测到 antiword，已跳过旧版 Word 纯文本抽取。")
        return ""
    try:
        completed = subprocess.run(
            [antiword, "-m", "UTF-8.txt", str(file_path)],
            check=False,
            capture_output=True,
            timeout=45,
        )
    except subprocess.TimeoutExpired:
        warnings.append("antiword 抽取超时，已尝试其他兼容方案。")
        return ""
    except Exception as exc:
        warnings.append(f"antiword 抽取异常: {_format_exception(exc)}")
        return ""

    stdout = completed.stdout.decode("utf-8", errors="ignore").strip()
    if completed.returncode != 0:
        stderr = completed.stderr.decode("utf-8", errors="ignore").strip()
        warnings.append(f"antiword 抽取失败: {stderr[:160] or '无法读取旧版 Word'}")
        return ""
    if stdout:
        warnings.append("已使用 antiword 抽取旧版 Word 正文。")
    return stdout


def _extract_legacy_doc_via_libreoffice(file_path: Path, warnings: list[str]) -> MaterialExtraction:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        warnings.append("未检测到 LibreOffice，无法将旧版 Word 转换为 docx/PDF 兜底。")
        return MaterialExtraction(method="legacy_doc_libreoffice_unavailable", source_kind="doc", warnings=warnings)
    try:
        with tempfile.TemporaryDirectory(prefix="material-ai-doc-convert-") as temp_dir:
            temp_path = Path(temp_dir)
            completed = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(temp_path),
                    str(file_path),
                ],
                check=False,
                capture_output=True,
                timeout=90,
            )
            if completed.returncode != 0:
                stderr = completed.stderr.decode("utf-8", errors="ignore").strip()
                warnings.append(f"LibreOffice 转 docx 失败: {stderr[:160] or '转换失败'}")
                return MaterialExtraction(method="legacy_doc_libreoffice_convert_failed", source_kind="doc", warnings=warnings)
            docx_files = sorted(temp_path.glob("*.docx"))
            if not docx_files:
                warnings.append("LibreOffice 转换未生成 docx，已尝试其他兜底。")
                return MaterialExtraction(method="legacy_doc_libreoffice_no_docx", source_kind="doc", warnings=warnings)
            extracted = _extract_docx(docx_files[0])
            extracted.method = "libreoffice_doc_to_docx_extract"
            if extracted.text.strip():
                warnings.append("已使用 LibreOffice 转换旧版 Word 后抽取正文。")
            return extracted
    except subprocess.TimeoutExpired:
        warnings.append("LibreOffice 转换旧版 Word 超时，已尝试其他兜底。")
    except Exception as exc:
        warnings.append(f"LibreOffice 转换旧版 Word 异常: {_format_exception(exc)}")
    return MaterialExtraction(method="legacy_doc_libreoffice_exception", source_kind="doc", warnings=warnings)


def _extract_xlsx(file_path: Path) -> MaterialExtraction:
    try:
        import openpyxl
    except ImportError as exc:
        raise HTTPException(500, f"缺少 Excel 解析依赖 openpyxl: {exc}")

    parts: list[str] = []
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    try:
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [_cell_to_text(value) for value in row]
                if any(values):
                    rows.append(values)
            if rows:
                parts.append(f"## 工作表：{sheet.title}\n\n{_rows_to_markdown(rows)}")
    finally:
        wb.close()

    text, truncated = _truncate_by_bytes("\n\n".join(parts), MAX_EXTRACT_TEXT_BYTES)
    warnings: list[str] = []
    images = _render_office_pages_to_images(file_path, ".xlsx", warnings) if len(text.strip()) < 800 else []
    return MaterialExtraction(
        text=text,
        method="openpyxl_tables",
        source_kind="xlsx",
        warnings=warnings,
        images=images,
        truncated=truncated,
    )


def _extract_xls(file_path: Path) -> MaterialExtraction:
    try:
        import xlrd
    except ImportError as exc:
        raise HTTPException(500, f"缺少 Excel 解析依赖 xlrd: {exc}")

    parts: list[str] = []
    wb = xlrd.open_workbook(str(file_path), formatting_info=False)
    for sheet in wb.sheets():
        rows = []
        for row_idx in range(sheet.nrows):
            values = [_cell_to_text(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
            if any(values):
                rows.append(values)
        if rows:
            parts.append(f"## 工作表：{sheet.name}\n\n{_rows_to_markdown(rows)}")

    text, truncated = _truncate_by_bytes("\n\n".join(parts), MAX_EXTRACT_TEXT_BYTES)
    warnings = ["旧版 Excel 已按单元格值抽取，合并单元格和图片需人工复核。"]
    images = _render_office_pages_to_images(file_path, ".xls", warnings) if len(text.strip()) < 800 else []
    return MaterialExtraction(
        text=text,
        method="xlrd_tables",
        source_kind="xls",
        warnings=warnings,
        images=images,
        truncated=truncated,
    )


def _extract_pdf(file_path: Path) -> MaterialExtraction:
    warnings: list[str] = []
    text = ""
    images: list[dict[str, str]] = []
    truncated = False
    if extract_document_text:
        extracted = extract_document_text(file_path, ".pdf", max_bytes=MAX_EXTRACT_TEXT_BYTES)
        text = extracted.text
        images = list(extracted.images or [])[:MAX_VISION_IMAGES]
        truncated = bool(extracted.truncated)
    if (not text.strip() or len(images) < 2) and render_pdf_pages_to_data_urls:
        rendered = render_pdf_pages_to_data_urls(file_path, dpi=144, max_pages=MAX_VISION_IMAGES)
        if rendered:
            images = rendered[:MAX_VISION_IMAGES]
            warnings.append("PDF 已渲染页面图像用于版式兜底。")
    return MaterialExtraction(
        text=text,
        method="pdf_text_and_render",
        source_kind="pdf",
        warnings=warnings,
        images=images,
        truncated=truncated,
    )


def _extract_image_file(file_path: Path, ext: str) -> MaterialExtraction:
    raw = file_path.read_bytes()
    if len(raw) > MAX_IMAGE_BYTES:
        raise HTTPException(413, "图片文件过大，无法直接用于 AI 视觉解析")
    mime = mimetypes.guess_type(file_path.name)[0] or ("image/jpeg" if ext in {".jpg", ".jpeg"} else "image/png")
    image_data = base64.b64encode(raw).decode("utf-8")
    return MaterialExtraction(
        text="",
        method="image_file",
        source_kind=ext.lstrip("."),
        images=[{"filename": file_path.name, "data_url": f"data:{mime};base64,{image_data}"}],
    )


def _render_office_pages_to_images(file_path: Path, ext: str, warnings: list[str]) -> list[dict[str, str]]:
    if not render_pdf_pages_to_data_urls:
        return []

    if ext not in {".doc", ".docx", ".xls", ".xlsx", ".ppt", ".pptx"}:
        return []

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        warnings.append("未检测到 LibreOffice，无法把该 Office 文档渲染为图片兜底。")
        return []

    try:
        with tempfile.TemporaryDirectory(prefix="material-ai-render-") as temp_dir:
            temp_path = Path(temp_dir)
            completed = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(temp_path),
                    str(file_path),
                ],
                check=False,
                capture_output=True,
                timeout=90,
            )
            if completed.returncode != 0:
                stderr = completed.stderr.decode("utf-8", errors="ignore").strip()
                warnings.append(f"Office 文档渲染失败: {stderr[:160] or 'LibreOffice 转换失败'}")
                return []
            pdf_files = sorted(temp_path.glob("*.pdf"))
            if not pdf_files:
                warnings.append("Office 文档渲染未生成 PDF，已跳过视觉兜底。")
                return []
            images = render_pdf_pages_to_data_urls(pdf_files[0], dpi=144, max_pages=MAX_VISION_IMAGES)
            if images:
                warnings.append("已将 Office 文档渲染为页面图片用于视觉兜底。")
            return images[:MAX_VISION_IMAGES]
    except subprocess.TimeoutExpired:
        warnings.append("Office 文档渲染超时，已跳过视觉兜底。")
    except Exception as exc:
        warnings.append(f"Office 文档渲染异常: {_format_exception(exc)}")
    return []


def _docx_table_to_markdown(table) -> str:
    return _rows_to_markdown(_docx_table_to_rows(table))


def _docx_table_to_rows(table) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in table._tbl.tr_lst:
        values: list[str] = []
        for cell in row.tc_lst:
            cell_text = "\n".join(
                text.strip()
                for text in cell.xpath(".//w:t/text()")
                if str(text or "").strip()
            )
            grid_span = 1
            tc_pr = cell.tcPr
            if tc_pr is not None and tc_pr.gridSpan is not None:
                try:
                    grid_span = max(1, int(tc_pr.gridSpan.val))
                except (TypeError, ValueError):
                    grid_span = 1
            values.append(cell_text)
            for _ in range(grid_span - 1):
                values.append("")
        if any(value.strip() for value in values):
            rows.append(values)
    return rows


def _extract_docx_zip_text(file_path: Path) -> str:
    try:
        with zipfile.ZipFile(file_path, "r") as zf:
            if "word/document.xml" not in zf.namelist():
                return ""
            xml_text = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    except Exception:
        return ""
    pieces = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml_text, flags=re.S)
    cleaned = [re.sub(r"<[^>]+>", "", piece).strip() for piece in pieces]
    return "\n".join(piece for piece in cleaned if piece)


def _rows_to_markdown(rows: list[list[str]]) -> str:
    normalized_rows = _normalize_table_rows(rows)
    if not normalized_rows:
        return ""
    header = normalized_rows[0]
    if len(normalized_rows) == 1:
        header = [f"列 {index + 1}" for index in range(len(normalized_rows[0]))]
        body_rows = normalized_rows
    else:
        body_rows = normalized_rows[1:]
    lines = [
        "| " + " | ".join(_escape_markdown_cell(value) for value in header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in body_rows:
        lines.append("| " + " | ".join(_escape_markdown_cell(value) for value in row) + " |")
    return "\n".join(lines)


def _normalize_table_rows(rows: list[list[str]]) -> list[list[str]]:
    if not rows:
        return []
    max_width = max((len(row) for row in rows), default=0)
    normalized = []
    for row in rows:
        values = [_cell_to_text(value) for value in row]
        values.extend([""] * (max_width - len(values)))
        if any(values):
            normalized.append(values)
    return normalized


def _csv_to_markdown(text: str) -> str:
    sample = text.splitlines()
    if not sample:
        return ""
    reader = csv.reader(sample)
    rows = [[cell.strip() for cell in row] for row in reader]
    return _rows_to_markdown(rows)


def _build_material_prompts(
    *,
    original_name: str,
    type_meta: dict[str, Any],
    extraction: MaterialExtraction,
    text_for_ai: str,
    vision_mode: bool,
) -> tuple[str, str]:
    schema_hint = _schema_hint(type_meta)
    system_prompt = (
        "你是广西外国语学院课程材料归档与格式化解析助手。"
        "你要根据真实教学归档材料的版式习惯，抽取可复用的结构化字段。"
        "必须只输出 JSON 对象，不要输出 Markdown 代码块、解释或寒暄。"
        "JSON 顶层字段必须包含 metadata, content_markdown, tables, warnings, export_payload。"
        "metadata 使用对象；content_markdown 使用完整 Markdown 正文；tables 使用数组；warnings 使用数组。"
        "export_payload 要保留后续按学校模板导出所需字段，至少包含 document_group, document_type, template_key, fields, sections, tables。"
    )
    sample_knowledge = _sample_knowledge(type_meta)
    if vision_mode:
        user_prompt = (
            f"请根据图片中的材料版式解析文件《{original_name}》。\n"
            f"材料大类：{type_meta['group_label']}；材料类型：{type_meta['label']}。\n"
            f"{sample_knowledge}\n\n"
            f"{schema_hint}\n\n"
            "如果图片中有手写签名、印章、勾选框、合并单元格或扫描痕迹，请尽量识别并在 warnings 中说明不确定项。"
        )
    else:
        user_prompt = (
            f"请解析文件《{original_name}》。\n"
            f"材料大类：{type_meta['group_label']}；材料类型：{type_meta['label']}。\n"
            f"本地抽取方式：{extraction.method}。\n"
            f"{sample_knowledge}\n\n"
            f"{schema_hint}\n\n"
            "以下是从原始文件抽取出的文本、表格或近似版式内容。请保留原文关键数据，不要编造缺失信息。\n\n"
            f"{text_for_ai}"
        )
    return system_prompt, user_prompt


def _schema_hint(type_meta: dict[str, Any]) -> str:
    key = type_meta["key"]
    common = (
        "通用字段建议：school, college, department, course_name, course_code, class_name, "
        "teacher_name, examiner_name, reviewer_name, leader_name, academic_year, semester, "
        "date, source_filename。所有能从表格标题、合并单元格、勾选项中读出的字段都要进入 metadata 和 export_payload.fields。"
    )
    specific = {
        "lesson_plan": "教案应抽取：课程性质、学分/学时、教材、章节、教学目标、重点难点、教学方法、教学过程、时间分配、课堂活动、板书/图示说明。",
        "teaching_calendar": "教学日历应抽取：课程代码、课程名称、班级、周次、日期、授课教师、学时、授课内容、教材章节、备注，并保留每一行进度。",
        "teaching_summary": "工作总结应抽取：院系、教师、课程、班级、日期、教学完成情况、学生学习情况、问题与改进、课程建设建议。",
        "evaluation_sheet": "评学表应抽取：课程、班级、教师、评价维度、分项结果、意见建议、统计信息。",
        "teaching_document": "教学文档应抽取：标题、适用课程、章节/主题、知识点、操作步骤、课堂任务、附件或图片说明。",
        "syllabus": "教学大纲应抽取：课程基本信息、课程性质、目标、内容模块、学时分配、教学方法、考核方式、教材与参考资料。",
        "assessment_plan": "考核计划表应抽取：assessment_type, assessment_method, assessment_items(assessment_form/content/score), total_score, examiner_name, reviewer_name, date, notes。",
        "grading_rubric": "评分细则应抽取：rubric_items(title/score/criteria), deduction_points, screenshot_requirements, examiner_name, reviewer_name, date, total_score，正文不能丢失扣分例外情况。",
        "exam_paper": "考核试卷应抽取：exam_flags, education_level, paper_type, exam_duration, score_table, paper_sections(title/score/content/tasks), student_fields, command_blocks, screenshot_requirements, total_score。",
        "final_teaching_summary": "教学工作总结应抽取：课程/班级/教师/日期、教学任务完成情况、考核与成绩分析、问题、改进建议。",
    }.get(key, "")
    final_schema = ""
    if key in FINAL_MATERIAL_TYPES:
        final_schema = (
            " export_payload.structured 必须包含与材料类型匹配的结构："
            "考核计划表包含 assessment_items；评分细则包含 rubric_items；试卷包含 paper_sections。"
            "表格请保留为 tables.rows，合并单元格不要重复造值，空占位可保留为空字符串。"
        )
    return f"{common}{specific}{final_schema}"


def _sample_knowledge(type_meta: dict[str, Any]) -> str:
    key = type_meta["key"]
    if key == "assessment_plan":
        return "样例特征：标题通常为“广西外国语学院课程考核计划表”，有学年学期、课程/班级/考核性质/教师/日期等表格字段，核心表格列为考核形式、考核技能或内容、分值。"
    if key == "grading_rubric":
        return "样例特征：标题通常为“广西外国语学院课程考核评分细则”，前部为课程元数据表，后部为评分细则正文或表格，签名可能以图片覆盖在表格附近。"
    if key == "exam_paper":
        return "样例特征：标题通常为“广西外国语学院课程考核试卷”，包含密封线、学生信息栏、考核方式勾选、成绩表和分题正文。"
    if key == "lesson_plan":
        return "样例特征：教案通常有封面页，包含课程名称、课程类别、学分学时、授课教师、使用教材，正文以周次/章节/目标/重点难点/教学过程表格组织。"
    if key == "teaching_calendar":
        return "样例特征：教学日历多为横向 Excel 表，表头包含课程代码、课程名称、授课班级、周学时、授课教师，正文按周次/日期/内容列出进度。"
    if key in {"teaching_summary", "final_teaching_summary"}:
        return "样例特征：工作总结为学校制式表格，顶部为学院、教师、课程、授课班级、日期等字段，主体是连续叙述文本。"
    return "样例特征：学校归档材料常使用制式标题、元数据表格、合并单元格和签名栏，请尽量保留字段和表格层次。"


def _build_local_fallback_result(
    *,
    original_name: str,
    type_meta: dict[str, Any],
    extraction: MaterialExtraction,
    warning: str,
) -> dict[str, Any]:
    metadata = {
        "source_filename": original_name,
        "document_group": type_meta["group_label"],
        "document_type": type_meta["label"],
        "title": Path(original_name).stem,
    }
    content = extraction.text.strip()
    return {
        "metadata": metadata,
        "content_markdown": content,
        "tables": [],
        "warnings": [warning],
        "export_payload": build_material_export_payload(
            document_group=type_meta["group_key"],
            document_type=type_meta["key"],
            type_meta=type_meta,
            metadata=metadata,
            content_markdown=content,
            tables=[],
        ),
    }


def _coerce_json_object(value: Any) -> dict[str, Any]:
    if isinstance(value, dict):
        return value
    text = str(value or "").strip()
    if not text:
        return {}
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start = text.find("{")
        end = text.rfind("}")
        if start >= 0 and end > start:
            try:
                return json.loads(text[start:end + 1])
            except json.JSONDecodeError:
                return {}
    return {}


def _coerce_dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _coerce_tables(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    tables = []
    for item in value:
        if isinstance(item, dict):
            tables.append(item)
        elif isinstance(item, list):
            tables.append({"title": "", "rows": item})
    return tables


def _merge_warnings(*values: Any) -> list[str]:
    merged: list[str] = []
    for value in values:
        if not value:
            continue
        if isinstance(value, str):
            candidates = [value]
        elif isinstance(value, list):
            candidates = [str(item) for item in value if item]
        else:
            candidates = [str(value)]
        for candidate in candidates:
            normalized = " ".join(candidate.split())
            if normalized and normalized not in merged:
                merged.append(normalized)
    return merged


def _normalize_academic_period(metadata: dict[str, Any]) -> None:
    text = " ".join(str(value) for value in metadata.values() if value)
    match = re.search(r"(20\d{2})\s*[-—~至]\s*(20\d{2})\s*学年", text)
    if match and not metadata.get("academic_year"):
        metadata["academic_year"] = f"{match.group(1)}-{match.group(2)}"
    semester_match = re.search(r"第\s*([一二三四五六七八九十1234567890]+)\s*学期", text)
    if semester_match and not metadata.get("semester"):
        metadata["semester"] = semester_match.group(1)


def _infer_export_sections(document_type: str, content_markdown: str) -> list[dict[str, Any]]:
    sections = []
    current_title = "正文"
    current_lines: list[str] = []
    for line in content_markdown.splitlines():
        heading = re.match(r"^(#{1,4})\s+(.+)$", line.strip())
        if heading:
            if current_lines:
                sections.append({"title": current_title, "content": "\n".join(current_lines).strip()})
                current_lines = []
            current_title = heading.group(2).strip()
        else:
            current_lines.append(line)
    if current_lines:
        sections.append({"title": current_title, "content": "\n".join(current_lines).strip()})
    if not sections:
        sections.append({"title": "正文", "content": content_markdown.strip()})
    return [{"type": document_type, **section} for section in sections if section.get("content")]


def _needs_vision_fallback(extraction: MaterialExtraction) -> bool:
    if extraction.images and not extraction.quality.get("usable", False):
        return True
    if extraction.images and not extraction.text.strip():
        return True
    if extraction.images and len(extraction.text.strip()) < 1200:
        return True
    return False


def _limit_text_for_ai(text: str) -> str:
    text = str(text or "").strip()
    if len(text) <= MAX_AI_TEXT_CHARS:
        return text
    return text[:MAX_AI_TEXT_CHARS] + "\n\n[内容过长，后续文本已截断]"


def _assess_text_quality(text: str, *, method: str = "") -> dict[str, Any]:
    sample = str(text or "")
    stripped = sample.strip()
    if not stripped:
        return {
            "status": "empty",
            "usable": False,
            "score": 0.0,
            "reasons": ["未抽取到可读文本。"],
        }

    quality_source = re.sub(
        r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$",
        "",
        stripped[:12000],
        flags=re.M,
    )
    compact = re.sub(r"\s+", "", quality_source)
    if len(compact) < MIN_USABLE_TEXT_CHARS:
        return {
            "status": "too_short",
            "usable": False,
            "score": 0.2,
            "reasons": ["抽取文本过短，无法作为可靠解析结果。"],
        }

    control_count = sum(1 for ch in compact if ord(ch) < 32 and ch not in "\t\n\r")
    replacement_count = compact.count("\ufffd")
    question_ratio = compact.count("?") / max(1, len(compact))
    cjk_count = sum(1 for ch in compact if "\u4e00" <= ch <= "\u9fff")
    ascii_word_count = sum(1 for ch in compact if ch.isascii() and (ch.isalnum() or ch.isspace()))
    common_punctuation = set("，。！？；：、“”‘’（）()[]{}《》<>-—_.,:/\\%+&=*@#|·~`'\" \n\r\t")
    readable_count = sum(
        1
        for ch in compact
        if ("\u4e00" <= ch <= "\u9fff")
        or ch.isalnum()
        or ch.isspace()
        or ch in common_punctuation
    )
    readable_ratio = readable_count / max(1, len(compact))
    cjk_ratio = cjk_count / max(1, len(compact))
    ascii_word_ratio = ascii_word_count / max(1, len(compact))
    long_symbol_runs = len(re.findall(r"[^\w\s\u4e00-\u9fff]{18,}", compact))
    binary_markers = any(marker in stripped[:20000] for marker in ("WordDocument", "Microsoft Word", "\x00", "\x01"))

    reasons: list[str] = []
    if control_count:
        reasons.append("文本包含不可见控制字符，疑似二进制内容。")
    if replacement_count:
        reasons.append("文本包含替换字符，疑似编码损坏。")
    if binary_markers:
        reasons.append("文本包含 Word 二进制结构标记。")
    if question_ratio > 0.08:
        reasons.append("文本中问号占比异常，疑似乱码。")
    if readable_ratio < 0.72:
        reasons.append("可读字符占比过低。")
    table_heavy_method = method in {"python_docx_tables", "openpyxl_tables", "xlrd_tables", "parsed_markdown"}
    if long_symbol_runs >= 2 and not (table_heavy_method and cjk_ratio >= 0.04):
        reasons.append("存在连续异常符号片段。")
    if len(compact) > 500 and cjk_ratio < 0.01 and ascii_word_ratio < 0.38:
        reasons.append("中文/英文正文特征不足，疑似从旧版 Office 中抽取到了底层结构。")

    status = "ok" if not reasons else "suspect"
    score = max(0.0, min(1.0, readable_ratio - (0.35 if reasons else 0.0)))
    usable = not reasons or (score >= 0.7 and not binary_markers and question_ratio <= 0.08)
    if method == "parsed_markdown" and len(compact) >= MIN_USABLE_TEXT_CHARS and readable_ratio >= 0.62 and not binary_markers:
        usable = True
        if not reasons:
            status = "ok"

    return {
        "status": status if usable else "failed",
        "usable": usable,
        "score": round(score, 3),
        "length": len(stripped),
        "readable_ratio": round(readable_ratio, 3),
        "cjk_ratio": round(cjk_ratio, 3),
        "question_ratio": round(question_ratio, 3),
        "method": method,
        "reasons": reasons,
    }


def _truncate_by_bytes(text: str, max_bytes: int) -> tuple[str, bool]:
    raw = str(text or "").encode("utf-8")
    if len(raw) <= max_bytes:
        return str(text or ""), False
    truncated = raw[:max_bytes].decode("utf-8", errors="ignore")
    return truncated, True


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return " ".join(str(value).replace("\r", "\n").split())


def _escape_markdown_cell(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", "<br>")


def _stringify_cell(value: Any) -> str:
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


def _strip_duplicate_title(content: str, title: str) -> str:
    lines = content.splitlines()
    while lines and not lines[0].strip():
        lines.pop(0)
    if not lines:
        return content
    first = lines[0].strip()
    normalized_first = first.lstrip("#").strip()
    if normalized_first == str(title or "").strip():
        return "\n".join(lines[1:]).strip()
    return content


def _metadata_key_label(key: str) -> str:
    labels = {
        "school": "学校",
        "college": "学院",
        "department": "系部",
        "course_name": "课程",
        "course_code": "课程代码",
        "class_name": "班级",
        "teacher_name": "教师",
        "academic_year": "学年",
        "semester": "学期",
        "date": "日期",
        "title": "标题",
    }
    return labels.get(key, key)


def _metadata_title(metadata: dict[str, Any], original_name: str, type_label: str) -> str:
    for key in ("title", "course_name"):
        value = str(metadata.get(key) or "").strip()
        if value:
            return value
    return f"{Path(original_name).stem}-{type_label}"


def _format_exception(exc: Exception) -> str:
    if isinstance(exc, HTTPException):
        return str(exc.detail)
    return str(exc)
