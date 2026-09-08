"""Fill the audited school's native Word table without rebuilding its layout.

Only editable cell contents and media are patched. In particular, saving through
python-docx must not rewrite styles, settings, VML, or any other original part.
Archived RTFs must have the audited layout before their values/chart can be used.
"""

from __future__ import annotations

from copy import deepcopy
import hashlib
import io
from pathlib import Path
import re
import unicodedata
from typing import Any
from zipfile import ZipFile

from .academic_final_material_source_service import (
    ACADEMIC_NATIVE_SOURCE_KEY, NativeAcademicSource, NativeAcademicSourceError,
)


_TEMPLATE_PATH = Path(__file__).with_name("assets") / "gxufl_academic_exam_analysis_template.docx"
_TEMPLATE_SHA256 = "073bbd7f624da1608919b20209e8aaa773aa945fb51eefffc16839cdf1d7d2b6"
_REMARK_ASSETS = {
    "已核": Path(__file__).with_name("assets") / "gxufl_exam_review_checked.png",
    "同意": Path(__file__).with_name("assets") / "gxufl_exam_review_agreed.png",
}
_IDENTITY_SLOTS = {"course_name": (3, 1), "course_hours": (3, 3), "department": (3, 5),
                   "teacher_name": (4, 1), "class_name": (7, 1)}


def _cell(table: Any, row: int, column: int) -> Any:
    from docx.table import _Cell
    return _Cell(table.rows[row]._tr.tc_lst[column], table)


def _set_value(cell: Any, value: Any, *, format_cell: Any = None) -> None:
    """Keep the source paragraph, run properties, and whitespace conventions."""
    from docx.oxml.ns import qn
    text = "" if value is None else str(value)
    nodes = list(cell._tc.iter(qn("w:t")))
    if nodes:
        nodes[0].text = text
        nodes[0].set(qn("xml:space"), "preserve")
        for node in nodes[1:]:
            node.text = ""
        return
    if not text:
        return
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    prototype = (format_cell or cell).paragraphs[0]
    properties = next((r._r.rPr for r in prototype.runs if r._r.rPr is not None), None)
    if properties is None and prototype._p.pPr is not None:
        properties = prototype._p.pPr.find(qn("w:rPr"))
    if properties is not None:
        run._r.insert(0, deepcopy(properties))


def _wrapped_lines(text: str, width: float, size: float, *, indent: float = 0) -> list[str]:
    """Conservative wrapping within a fixed school-form slot, never shrink text.

    CJK glyphs use an em; Latin uses at least 0.7 em (wide letters a full em).
    Used for a conservative capacity check; narrative text retains native Word
    wrapping so a decimal or percentage is never split by an inserted break.
    """
    lines, current = [], ""
    used = indent
    for char in text:
        advance = size if unicodedata.east_asian_width(char) in {"W", "F"} or char in "MW@%" else size * .7
        if current and used + advance > width:
            lines.append(current)
            current, used = "", 0
        current += char
        used += advance
    return lines + [current]


def _fill_analysis(cell: Any, value: Any) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt

    text = str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not text:
        return
    paragraphs = text.split("\n")
    wrapped = []
    for line in paragraphs:
        heading = bool(re.match(r"^[一二三四五六七八九十]+[、.]", line))
        indent = 0 if heading or not line else 24
        wrapped.append((_wrapped_lines(line, cell.width.pt - 3, 12, indent=indent), indent))
    line_count = sum(len(lines) for lines, _ in wrapped)
    # The source has 142 pt, fixed. Leave room for Word's cell/paragraph marks.
    if line_count * 16.2 > 140:
        raise ValueError(
            f"分析正文按原版字号排版需要约 {line_count} 行，超过原版单元格的 8 行容量。"
            "请精简正文后导出；已保存内容不会被裁切，也不会缩小字体或改变表格尺寸。"
        )
    prototype = deepcopy(cell.paragraphs[0]._p)
    for paragraph in list(cell._tc.findall(qn("w:p"))):
        cell._tc.remove(paragraph)
    from docx.text.paragraph import Paragraph
    for line, (_, indent) in zip(paragraphs, wrapped):
        element = deepcopy(prototype)
        for child in list(element):
            if child.tag != qn("w:pPr"):
                element.remove(child)
        cell._tc.append(element)
        paragraph = Paragraph(element, cell)
        paragraph.paragraph_format.first_line_indent = Pt(indent)
        # Inherit the source's single spacing; a fixed 14.4 pt line would
        # compress the original Song/Times mixed text (about 16.2 pt in Word).
        paragraph.paragraph_format.widow_control = False
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        run.add_text(line)


def _overlay_image(paragraph: Any, path_value: Any, *, x: float, y: float,
                   width: float, height: float, description: str) -> None:
    """Floating images occupy the original blank area; they cannot resize rows."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from PIL import Image

    if not path_value:
        return
    path = Path(str(path_value))
    if not path.is_file():
        raise ValueError("审核图片已不可用，请重新绑定签名后导出。")
    with Image.open(path) as picture:
        scale = min(width / picture.width, height / picture.height)
        actual_width, actual_height = picture.width * scale, picture.height * scale
    shape = paragraph.add_run().add_picture(str(path), width=Pt(actual_width), height=Pt(actual_height))
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {"distT": "0", "distB": "0", "distL": "0", "distR": "0",
                        "simplePos": "0", "relativeHeight": "251659264", "behindDoc": "0",
                        "locked": "0", "layoutInCell": "0", "allowOverlap": "1"}.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for direction, coordinate in (("H", x + (width - actual_width) / 2), ("V", y + (height - actual_height) / 2)):
        position = OxmlElement(f"wp:position{direction}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = str(int(Pt(coordinate)))
        position.append(offset)
        anchor.append(position)
    for name in ("extent", "effectExtent"):
        node = inline.find(qn(f"wp:{name}"))
        if node is not None:
            anchor.append(deepcopy(node))
    anchor.append(OxmlElement("wp:wrapNone"))
    for name in ("docPr", "cNvGraphicFramePr"):
        node = deepcopy(inline.find(qn(f"wp:{name}")))
        if name == "docPr":
            node.set("descr", description)
        anchor.append(node)
    anchor.append(deepcopy(inline.find(qn("a:graphic"))))
    inline.getparent().replace(inline, anchor)


def _fill_reviews(document: Any, table: Any, fields: dict[str, Any]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt

    # The source starts after its retained 1 pt paragraph (1.2 pt line height).
    table_y = document.sections[0].top_margin.pt + 1.2
    review_y = table_y + sum(row.height.pt for row in table.rows[:20])
    left = document.sections[0].left_margin.pt + .25  # Native tblInd=5 twips.
    for column, role in enumerate(("department", "dean")):
        cell = _cell(table, 20, column)
        cell_width = cell.width.pt
        paragraph = cell.paragraphs[0]
        key = f"{role}_review_opinion"
        value = str(fields.get(key) or "").strip()
        if len(value) > 80:
            raise ValueError("审核意见不能超过 80 字，请精简后导出。")
        remark_path = _REMARK_ASSETS.get(value) if value else (fields.get(f"{role}_review_opinion_image_path") if key not in fields else None)
        if remark_path:
            _overlay_image(paragraph, remark_path, x=left + (cell_width - 62) / 2,
                           y=review_y + 4, width=62, height=32, description=value or "审核批语")
        elif value:
            # An explicit opinion replaces the blank slot only. Its paragraph
            # remains in the original 93 pt row, with room for the signature.
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.left_indent = Pt(2)
            paragraph.paragraph_format.right_indent = Pt(2)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(13.2)
            lines = _wrapped_lines(value, cell_width - 5, 11)
            if len(lines) > 4:
                raise ValueError("审核意见超出原版填写区域，请精简后导出。")
            run = paragraph.add_run("\n".join(lines))
            run.font.name, run.font.size = "宋体", Pt(11)
            run._r.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
        ids = fields.get(f"{role}_personal_signature_ids", fields.get(f"{role}_signature_ids"))
        count = len(ids) if isinstance(ids, list) else 1
        signature_width = min(cell_width - 12, 138 if count > 1 else 88)
        _overlay_image(paragraph, fields.get(f"{role}_signature_image_path"),
                       x=left + cell_width - signature_width - 6, y=review_y + 59,
                       width=signature_width, height=30, description="审核签名")
        left += cell_width


def _preserving_package(original: bytes, document: Any, media: dict[str, bytes]) -> bytes:
    """Copy every untouched ZIP member byte-for-byte, including VML and styles."""
    serialized = io.BytesIO()
    document.save(serialized)
    result = io.BytesIO()
    editable_parts = {"word/document.xml", "word/_rels/document.xml.rels", "[Content_Types].xml"}
    with ZipFile(io.BytesIO(original)) as baseline, ZipFile(serialized) as edited, ZipFile(result, "w") as target:
        for entry in baseline.infolist():
            data = media.get(entry.filename)
            if data is None:
                data = edited.read(entry.filename) if entry.filename in editable_parts else baseline.read(entry.filename)
            target.writestr(entry, data)
        for entry in edited.infolist():
            if entry.filename not in baseline.namelist():
                target.writestr(entry, edited.read(entry.filename))
    return result.getvalue()


def build_native_exam_analysis_docx(parse_payload: dict[str, Any]) -> bytes:
    from docx import Document
    from .academic_final_material_document_service import _academic_period_text, _chart_image, _score_text
    from .academic_final_material_service import parse_exam_analysis_rtf

    try:
        original = _TEMPLATE_PATH.read_bytes()
    except OSError as exc:
        raise NativeAcademicSourceError("试卷分析表原生模板不可用，请修复模板后导出。") from exc
    if hashlib.sha256(original).hexdigest() != _TEMPLATE_SHA256:
        raise NativeAcademicSourceError("试卷分析表原生模板校验失败，请修复模板后导出。")
    payload = parse_payload.get("export_payload", parse_payload)
    fields = dict(payload.get("fields") or {})
    structured = dict(payload.get("structured") or {})
    media = {}
    source = parse_payload.get(ACADEMIC_NATIVE_SOURCE_KEY)
    if source is not None:
        if not isinstance(source, NativeAcademicSource) or hashlib.sha256(source.content).hexdigest() != source.sha256:
            raise NativeAcademicSourceError("教务原件校验失败，请重新同步后导出。")
        from .academic_exam_analysis_rtf_service import inspect_native_exam_analysis_source
        inspected = inspect_native_exam_analysis_source(source.content)
        parsed = parse_exam_analysis_rtf(source.content)
        # Identity, statistics and the chart all come from the SAME original.
        # Editable choices, narrative and authorized signatures stay in fields.
        for key in (*_IDENTITY_SLOTS, "academic_year", "semester"):
            fields[key] = parsed["fields"].get(key, "")
        supplied = structured.get("score_distribution")
        if isinstance(supplied, list) and supplied and [int(item.get("count") or 0) for item in supplied] != [item["count"] for item in parsed["distribution"]]:
            raise NativeAcademicSourceError("已保存成绩分布与教务原图不一致，请重新同步后导出。")
        supplied_statistics = structured.get("statistics") or {}
        for key, value in parsed["statistics"].items():
            saved = supplied_statistics.get(key)
            if saved is not None and value is not None and abs(float(saved) - value) > .02:
                raise NativeAcademicSourceError("已保存成绩统计与教务原件不一致，请重新同步后导出。")
        structured["score_distribution"] = parsed["distribution"]
        structured["statistics"] = parsed["statistics"]
        media = {"word/media/image1.png": inspected["vertical_label_png"],
                 "word/media/image2.png": inspected["chart_png"],
                 "word/media/image3.png": inspected["analysis_label_png"]}
    document = Document(io.BytesIO(original))
    table = document.tables[0]
    if source is not None:
        _set_value(_cell(table, 2, 0), inspected["hidden_text"])
    _set_value(_cell(table, 1, 0), _academic_period_text(fields))
    for key, (row, column) in _IDENTITY_SLOTS.items():
        _set_value(_cell(table, row, column), fields.get(key, ""))
    for key, row, choices in (
        ("course_nature", 4, ((4, "选修"), (6, "必修"))),
        ("proposition_form", 5, ((2, "试题库"), (4, "试卷库"), (6, "教师组题"))),
        ("exam_form", 6, ((2, "开卷"), (4, "闭卷"))),
        ("separate_teaching_exam", 6, ((7, "是"), (9, "否"))),
    ):
        for column, label in choices:
            _set_value(_cell(table, row, column), "√" if fields.get(key) == label else "", format_cell=_cell(table, row, column - 1))
    for column, label in enumerate(("本人阅卷", "同行阅卷", "集体阅卷", "机器阅卷", "其他"), 1):
        if fields.get("marking_form") == label:
            _set_value(_cell(table, 14, column), label + " √")
    distribution = structured.get("score_distribution") or []
    for index in range(5):
        item = distribution[index] if index < len(distribution) else {}
        _set_value(_cell(table, 10, index + 2), int(item.get("count") or 0))
        _set_value(_cell(table, 11, index + 2), f"{float(item.get('ratio') or 0):.2f}%")
    statistics = structured.get("statistics") or {}
    for row, column, key, format_string in ((12, 2, "average", "{:.2f}"), (12, 4, "standard_deviation", "{:.2f}"),
                                           (13, 2, "maximum", None), (13, 4, "minimum", None), (13, 6, "pass_rate", "{:.2f}%")):
        value = statistics.get(key)
        _set_value(_cell(table, row, column), format_string.format(float(value or 0)) if format_string else _score_text(value))
    if not media:
        media["word/media/image2.png"] = _chart_image(distribution).getvalue()
    _fill_analysis(_cell(table, 18, 1), structured.get("analysis_text", fields.get("analysis_text", "")))
    _fill_reviews(document, table, fields)
    return _preserving_package(original, document, media)
